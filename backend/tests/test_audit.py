"""Audit tests (TN-only): category completeness, formulas, traceability, no silent ₹0, land separation, locking."""
from __future__ import annotations

import io

import pytest

from conftest import new_project

CATS = ["Doors & Windows", "Earthwork", "Concrete Works", "RCC Works", "Reinforcement", "Masonry", "Flooring", "Plastering", "Painting"]


@pytest.fixture(scope="module")
def proj(env):
    _, c = env
    return new_project(c, dict(district="Chennai", city="Chennai Corporation"))


def items(p):
    return {i["key"]: i for i in p["estimate"]["items"]}


def test_every_category_quantified_and_either_priced_or_listed(proj):
    """Every client category has quantities; each line is priced from TN data or listed with NEEDS RATE / NEEDS RATE ANALYSIS."""
    ex = {x["key"] for x in proj["estimate"]["excluded"]}
    for c in CATS:
        lines = [i for i in proj["estimate"]["items"] if i["category"] == c and i["qty"] > 0]
        assert lines, f"{c} not quantified"
        for i in lines:
            assert (i["rate"] is not None and i["amount"] > 0) or (i["key"] in ex and i["rate_gap"] in ("NEEDS RATE", "NEEDS RATE ANALYSIS"))
    priced = {i["category"] for i in proj["estimate"]["items"] if i["rate"] is not None}
    assert {"Earthwork", "Concrete Works", "RCC Works", "Reinforcement", "Masonry", "Flooring", "Plastering", "Painting"} <= priced


def test_doors_windows_counts_sizes_basis(proj):
    ops = [o for s in proj["sheets"] for o in s.get("openings", [])]
    assert any(o["kind"] == "door" for o in ops) and any(o["kind"] == "window" for o in ops)
    assert all(o["w"] > 0.4 for o in ops)
    dw = [i for i in proj["estimate"]["items"] if i["category"] == "Doors & Windows"]
    assert dw and all(i["basis"] for i in dw)
    # PWD data gaps for doors / windows: an editable estimated average rate, with the PWD suggestion still offered
    sh = items(proj)["door_shutter"]
    assert sh["map_method"] == "estimated" and sh["rate"] > 0 and sh["derived"]["basis"]
    assert "M-0387" in (sh.get("suggested") or [c["code"] for c in sh.get("candidates", [])])


def test_quantity_formulas_and_amounts(proj):
    it = items(proj)
    carpet = sum(r["area"] for s in proj["sheets"] if s.get("kind") == "floor_plan" for r in s["rooms"])
    assert abs(it["pcc_floor"]["qty"] - round(carpet * 0.075, 2)) < 0.05 * max(1, it["pcc_floor"]["qty"])
    for i in proj["estimate"]["items"]:
        if i["rate"] is not None:
            assert abs(i["amount"] - i["qty"] * i["rate"]) <= 0.01 * max(1, i["amount"]) + 1
    e = proj["estimate"]
    assert abs(e["total"] - (e["subtotal"] + e["contingency"] + e["cess"] + e["gst"])) < 1


def test_no_silent_zero(proj):
    for i in proj["estimate"]["items"]:
        assert not (i["rate"] is not None and i["qty"] > 0 and i["amount"] == 0)
    unpriced = [i for i in proj["estimate"]["items"] if i["rate"] is None]
    assert len(proj["estimate"]["excluded"]) == len(unpriced) and proj["estimate"]["complete"] is (not unpriced)
    assert all(x["reason"] for x in proj["estimate"]["excluded"])
    est = [i for i in proj["estimate"]["items"] if i.get("map_method") == "estimated"]
    assert len(proj["estimate"]["estimated"]) == len(est) and all(x["basis"] for x in proj["estimate"]["estimated"])


def test_report_traceability(env, proj):
    _, c = env
    for i in proj["estimate"]["items"]:
        assert i["formula"] and i.get("basis_class")
        if i["rate"] is not None:
            pv = i["rate_provenance"]
            if i.get("map_method") == "estimated":
                assert pv["kind"] == "Estimated average rate" and pv["basis"]
            else:
                assert pv and pv["sha256"] and pv["page"] and i["rate_schedule"]
    html = c.get(f"/api/projects/{proj['id']}/report").text
    assert "estimated average rates" in html and "Tamil Nadu" not in html  # client-facing report is de-branded
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(c.get(f"/api/projects/{proj['id']}/export.xlsx").content))
    hdr = list(next(wb["Detailed BOQ"].iter_rows(values_only=True)))
    for h in ("Rate schedule", "Authority", "Year / revision", "Rate type", "Basis class"):
        assert h in hdr
    assert "Not in total" in wb.sheetnames


def test_other_elements_detected_not_ignored():
    import takeoff
    sh = [dict(file="a.pdf", sheet="1", texts=["LIFT", "RAMP 1:12", "SUMP 10000 L"], rooms=[dict(name="Sunshade")])]
    names = {o["element"] for o in takeoff.other_elements(sh)}
    assert {"Lift / elevator", "Ramp", "Sump / underground tank", "Sunshade / chajja"} <= names


def test_land_separate_and_persisted(env, proj):
    _, c = env
    C = proj["estimate"]["total"]
    r = c.patch(f"/api/projects/{proj['id']}", json=dict(params=dict(land_area=100.0, land_rate=50000.0, stamp_duty_pct=0.0))).json()
    assert r["estimate"]["total"] == C and abs(r["land"]["total_with_land"] - (C + 5_000_000)) < 1
    r = c.patch(f"/api/projects/{proj['id']}", json=dict(params=dict(land_rate=60000.0))).json()
    assert r["estimate"]["total"] == C and abs(r["land"]["total_with_land"] - (C + 6_000_000)) < 1
    g = c.get(f"/api/projects/{proj['id']}").json()
    assert g["land"]["land_cost"] == 6_000_000.0 and g["estimate"]["total"] == C


def test_override_persists_and_locking(env):
    _, c = env
    p = new_project(c, {})
    r = c.patch(f"/api/projects/{p['id']}/items/door_shutter", json=dict(rate=2500.0))
    assert r.status_code == 200
    g = c.get(f"/api/projects/{p['id']}").json()
    d = next(i for i in g["estimate"]["items"] if i["key"] == "door_shutter")
    assert d["rate"] is not None and d["amount"] > 0 and not any(x["key"] == "door_shutter" for x in g["estimate"]["excluded"])
    assert c.post(f"/api/projects/{p['id']}/confirm", json=dict(by="test")).status_code == 200
    assert c.patch(f"/api/projects/{p['id']}/items/door_shutter", json=dict(rate=1.0)).status_code == 409
    assert c.post(f"/api/projects/{p['id']}/reopen", json={}).status_code == 200
