"""Flexible estimation scope — acceptance tests against the real API, parsers, take-off and pricing (no mocks).

Drawings: generated sample sets (data/samples) and a single-room DXF drawn here with the same CAD conventions.
"""
from __future__ import annotations

import io
import json
import os
import tempfile

import pytest

ROOM_DXF = None


def room_dxf() -> bytes:
    """One bedroom 3.6 × 4.2 m clear, 230 mm walls, one door, one window, label + written dimensions."""
    global ROOM_DXF
    if ROOM_DXF is None:
        import samples
        t, W, D = 0.23, 3.6 + 0.23, 4.2 + 0.23
        walls = [dict(x1=0, y1=0, x2=W, y2=0, t=t, ext=True), dict(x1=W, y1=0, x2=W, y2=D, t=t, ext=True),
                 dict(x1=W, y1=D, x2=0, y2=D, t=t, ext=True), dict(x1=0, y1=D, x2=0, y2=0, t=t, ext=True)]
        plan = dict(walls=walls, columns=[],
                    openings=[dict(kind="door", x=1.0, y=0, w=0.9, horiz=True), dict(kind="window", x=W / 2, y=D, w=1.5, horiz=True)],
                    rooms=[dict(name="BEDROOM", poly=[(t / 2, t / 2), (W - t / 2, t / 2), (W - t / 2, D - t / 2), (t / 2, D - t / 2)], clear=(3.6, 4.2))])
        p = os.path.join(tempfile.mkdtemp(), "bedroom.dxf")
        samples.render_dxf(plan, p)
        ROOM_DXF = open(p, "rb").read()
    return ROOM_DXF


def sample(A, name) -> bytes:
    return open(os.path.join(A.SAMPLES, name), "rb").read()


def new(c, name, scope=None, **kw):
    data = dict(name=name, **kw)
    if scope:
        data["scope"] = json.dumps(scope)
    r = c.post("/api/projects", data=data)
    assert r.status_code == 200, r.text
    return r.json()


def upload(c, pid, *files):
    r = c.post(f"/api/projects/{pid}/files", files=[("files", (n, b)) for n, b in files])
    assert r.status_code == 200, r.text
    return r.json()


def items(p):
    return {i["key"]: i for i in p["estimate"]["items"]}


def reconciles(e):
    inscope = round(sum(i["amount"] for i in e["items"] if i["in_scope"]), 2)
    cats = round(sum(x["amount"] for x in e["categories"]), 2)
    return abs(inscope - e["subtotal"]) < 0.05 and abs(cats - e["subtotal"]) < 0.05 and \
        abs(e["total"] - (e["subtotal"] + e["contingency"] + e["cess"] + e["gst"])) < 0.05


# 1, 2 ------------------------------------------------------------------------------------------------------------
def test_single_room_estimate_without_floor_information(env):
    _, c = env
    p = new(c, "Bedroom", dict(type="room", label="Master bedroom"))
    assert p["sheets"] == []  # created with no drawings and no building configuration
    q = upload(c, p["id"], ("bedroom.dxf", room_dxf()))
    plans = [s for s in q["sheets"] if s["kind"] == "floor_plan"]
    assert plans and any("BED" in r["name"].upper() for r in plans[0]["rooms"])
    assert q["estimate"]["total"] > 0 and reconciles(q["estimate"])
    assert not any(x["code"] in ("STOREY COUNT REQUIRED", "FLOOR COUNT CONFLICT") for x in q["checks"])
    cats = {x["name"] for x in q["estimate"]["categories"]}
    assert "RCC Works" not in cats and "Earthwork" not in cats  # structure / substructure outside a room scope
    assert {x["category"] for x in q["estimate"]["out_of_scope"]} >= {"Earthwork", "RCC Works"}
    # even if a storey count is set, a room scope never multiplies by it; the whole-property scope does
    c.patch(f"/api/projects/{p['id']}", json=dict(params=dict(floors=3)))
    room = c.get(f"/api/projects/{p['id']}").json()
    prop = c.patch(f"/api/projects/{p['id']}", json=dict(scope=dict(type="property"))).json()
    assert abs(items(prop)["plaster_int"]["qty"] - 3 * items(room)["plaster_int"]["qty"]) <= 0.02 * items(prop)["plaster_int"]["qty"]


# 3 ---------------------------------------------------------------------------------------------------------------
def test_section_scope_and_explicit_repeat(env):
    _, c = env
    p = new(c, "Wing B", dict(type="section", label="Wing B", config=dict(repeat=1)))
    q1 = upload(c, p["id"], ("bedroom.dxf", room_dxf()))
    q2 = c.patch(f"/api/projects/{p['id']}", json=dict(scope=dict(type="section", label="Wing B", config=dict(repeat=3)))).json()
    a, b = items(q1)["plaster_int"]["qty"], items(q2)["plaster_int"]["qty"]
    assert abs(b - 3 * a) <= 0.02 * b  # repeated only because the user said "3"
    assert q2["scope"]["config"]["repeat"] == "3" or q2["scope"]["config"]["repeat"] == 3


# 4, 5, 6 ---------------------------------------------------------------------------------------------------------
def test_floors_different_layouts_and_typical_repeat_without_overcount(env):
    A, c = env
    p = new(c, "Two floors", dict(type="floors"))
    q = upload(c, p["id"], ("villa_g1.dxf", sample(A, "villa_g1.dxf")))
    assert q["estimate"]["total"] > 0  # 4: a single floor
    i0 = {s["sheet"] if False else n: s for n, s in enumerate(q["sheets"])}
    idx = next(n for n, s in i0.items() if s["kind"] == "floor_plan")
    one = c.patch(f"/api/projects/{p['id']}/sheets/{idx}", json=dict(floor=dict(label="Ground floor", level=0, repeat=1))).json()
    # 5: a different layout on the first floor
    q2 = upload(c, p["id"], ("bedroom_ff.dxf", room_dxf()))
    idx2 = max(n for n, s in enumerate(q2["sheets"]) if s["kind"] == "floor_plan" and s["file"] == "bedroom_ff.dxf")
    two = c.patch(f"/api/projects/{p['id']}/sheets/{idx2}", json=dict(floor=dict(label="First floor", level=1, repeat=1))).json()
    assert items(two)["plaster_int"]["qty"] > items(one)["plaster_int"]["qty"]
    # 6: typical floor repeated ×3 — finishes ×3 for that plan, foundation NOT ×3
    c.delete(f"/api/projects/{p['id']}/files/bedroom_ff.dxf")
    base = c.patch(f"/api/projects/{p['id']}/sheets/{idx}", json=dict(floor=dict(label="Ground floor", level=0, repeat=1))).json()
    rep3 = c.patch(f"/api/projects/{p['id']}/sheets/{idx}", json=dict(floor=dict(label="Typical floor", level=0, repeat=3))).json()
    assert rep3["params"]["floors"] == 3
    assert abs(items(rep3)["plaster_int"]["qty"] - 3 * items(base)["plaster_int"]["qty"]) <= 0.05 * items(rep3)["plaster_int"]["qty"]
    for k in ("excavation", "rcc_footing", "pcc_found"):
        if k in items(base):
            assert items(rep3)[k]["qty"] <= items(base)[k]["qty"] * 1.0001, k  # one foundation, whatever the storeys


# 7 ---------------------------------------------------------------------------------------------------------------
def test_entire_property_many_drawings(env):
    A, c = env
    p = new(c, "Villa", dict(type="property"))
    q = upload(c, p["id"], ("villa_g1.dxf", sample(A, "villa_g1.dxf")), ("drawing_set.pdf", sample(A, "drawing_set.pdf")))
    assert len({s["file"] for s in q["sheets"]}) == 2 and q["estimate"]["total"] > 0
    assert q["estimate"]["out_of_scope"] == [] and reconciles(q["estimate"])


# 8 ---------------------------------------------------------------------------------------------------------------
def test_custom_scope_selected_categories_only(env):
    A, c = env
    p = new(c, "Repaint + floor", dict(type="custom", categories=["Painting", "Flooring"]), sample="compact_home.dxf")
    assert {x["name"] for x in p["estimate"]["categories"]} <= {"Painting", "Flooring"}
    assert p["estimate"]["out_of_scope"] and all(x["category"] not in ("Painting", "Flooring") for x in p["estimate"]["out_of_scope"])
    assert reconciles(p["estimate"])


# 9, 10 -----------------------------------------------------------------------------------------------------------
def test_scale_and_structural_gaps_are_flagged(env):
    _, c = env
    p = new(c, "Photo plan", None, sample="villa_g1.png")
    codes = {x["code"] for x in p["checks"]}
    assert codes & {"SCALE ASSUMED", "SCALE UNRESOLVED"}
    assert "STRUCTURAL INPUT MISSING" in codes


# 11 --------------------------------------------------------------------------------------------------------------
def test_duplicate_drawings_not_silently_double_counted(env):
    A, c = env
    p = new(c, "Dup", None, sample="villa_g1.dxf")
    t0 = p["estimate"]["total"]
    q = upload(c, p["id"], ("copy of villa.dxf", sample(A, "villa_g1.dxf")))  # identical bytes
    assert q["upload_skipped"] and len(q["sheets"]) == len(p["sheets"]) and q["estimate"]["total"] == t0
    q2 = upload(c, p["id"], ("villa_g1.pdf", sample(A, "villa_g1.pdf")))  # same floor, other format
    assert any(x["code"] in ("SAME FLOOR IN TWO DRAWINGS", "POSSIBLE DUPLICATE DRAWING") for x in q2["checks"])
    assert q2["estimate"]["total"] <= t0 * 1.05  # not doubled
    # a different file under an existing name is kept separately, never overwrites
    q3 = upload(c, p["id"], ("villa_g1.dxf", room_dxf()))
    assert any(x["file"].startswith("villa_g1 (") for x in c.get(f"/api/projects/{p['id']}/uploads").json())


# 12, 13, 14, 15 --------------------------------------------------------------------------------------------------
def test_corrections_persist_through_reload_new_drawings_and_removal(env):
    A, c = env
    p = new(c, "Persist", None, sample="compact_home.dxf")
    pid = p["id"]
    k = next(i["key"] for i in p["estimate"]["items"] if i["key"] == "plaster_int")
    est = next(i for i in p["estimate"]["items"] if i.get("map_method") == "estimated")
    assert c.patch(f"/api/projects/{pid}/items/{k}", json=dict(qty=123.45)).status_code == 200
    assert c.patch(f"/api/projects/{pid}/items/{est['key']}", json=dict(rate=777.0)).status_code == 200
    r = c.get(f"/api/projects/{pid}").json()  # 12/13: reopened from the database
    assert items(r)[k]["qty"] == 123.45 and items(r)[est["key"]]["base_rate"] == 777.0
    assert items(r)[est["key"]]["map_method"] == "estimated"  # still labelled as an estimate-origin line
    other = next(i for i in r["estimate"]["items"] if i.get("map_method") == "estimated" and i["key"] != est["key"])
    assert other["rate_provenance"]["kind"] == "Estimated average rate"
    r2 = upload(c, pid, ("bedroom.dxf", room_dxf()))  # 14: new drawing, edits survive
    assert items(r2)[k]["qty"] == 123.45 and items(r2)[est["key"]]["base_rate"] == 777.0
    n = len(r2["sheets"])
    r3 = c.delete(f"/api/projects/{pid}/files/bedroom.dxf").json()  # 15: removal
    assert len(r3["sheets"]) < n and all(s["file"] != "bedroom.dxf" for s in r3["sheets"])
    assert "bedroom.dxf" not in [x["file"] for x in c.get(f"/api/projects/{pid}/uploads").json()]
    assert items(r3)[k]["qty"] == 123.45
    assert c.delete(f"/api/projects/{pid}/files/nothing.dxf").status_code == 404


# 16, 17, 18 ------------------------------------------------------------------------------------------------------
def test_partial_estimate_report_reconciliation_and_exports(env):
    _, c = env
    p = new(c, "Partial", dict(type="custom", categories=["Painting"]), sample="compact_home.dxf")
    html = c.get(f"/api/projects/{p['id']}/report").text
    assert "PARTIAL estimate" in html and "outside the selected scope" in html
    e = p["estimate"]
    assert reconciles(e)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(c.get(f"/api/projects/{p['id']}/export.xlsx").content))
    rows = list(wb["Not in total"].iter_rows(values_only=True))
    assert any("outside the selected estimation scope" in str(r[-1]) for r in rows)
    pdf = c.get(f"/api/projects/{p['id']}/report.pdf")
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
    for fmt in ("compact_home.dxf", "compact_home.pdf", "compact_home.png"):  # supported formats still work
        assert new(c, fmt, None, sample=fmt)["sheets"]


def test_scope_options_and_invalid_scope(env):
    _, c = env
    o = c.get("/api/scope-options").json()
    assert set(o["types"]) == {"room", "section", "floors", "property", "custom"} and "Bedroom" in o["labels"]
    assert c.post("/api/projects", data=dict(name="x", scope="not json")).status_code == 400
    p = new(c, "Unknown type", dict(type="galaxy"))
    assert p["scope"]["type"] == "property"  # unknown → safe default, never an error-prone guess
