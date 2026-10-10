"""Drawing-based PWD cost estimation system (FastAPI backend)."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import shutil
import sqlite3
import threading
import time
import uuid
from datetime import datetime

import cv2
import numpy as np
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles

import floors
import parsers
import rates
import rooms as rooms_mod
import areas
import learning
import openings
import quality
import tn_pricing
import tn_scope
import samples
import takeoff
import scope as scope_mod
from ml import Models
from sor import CATEGORY_ORDER, SEED_ANALYSIS, SEED_SOR, analysis_rate

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _load_dotenv(path):
    """KEY=VALUE lines from the git-ignored .env (API keys); real environment variables win."""
    try:
        for line in open(path, encoding="utf-8"):
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))
    except OSError:
        pass


if not os.environ.get("EST_NO_DOTENV"):
    _load_dotenv(os.path.join(ROOT, ".env"))
REPO_DATA = os.path.join(ROOT, "data")  # models, reference data, the CPWD DSR source file, samples
DATA = os.environ.get("EST_DATA_DIR") or REPO_DATA  # database, uploads, previews (overridable, e.g. for tests)
UPLOADS = os.path.join(DATA, "uploads")
PREVIEWS = os.path.join(DATA, "previews")
SAMPLES = os.path.join(REPO_DATA, "samples")
FRONT = os.path.join(ROOT, "frontend")
DB = os.path.join(DATA, "estimate.db")
JURIS = os.path.join(REPO_DATA, "jurisdictions.json")
for d in (DATA, UPLOADS, PREVIEWS, SAMPLES):
    os.makedirs(d, exist_ok=True)

app = FastAPI(title="Tamil Nadu PWD Drawing-Based Cost Estimation", version="3.0")
LOCK = threading.RLock()
STATE: dict = {}
# plausible ranges for user-entered parameters (a value outside is rejected with 400, never silently used)
PARAM_BOUNDS = dict(floors=(1, 200, int), floor_height=(2.0, 10.0, float), plinth_height=(0.0, 3.0, float),
                    foundation_depth=(0.3, 10.0, float), slab_thickness=(0.075, 0.6, float), parapet_height=(0.0, 3.0, float),
                    sill_height=(0.0, 3.0, float), lintel_height=(1.5, 5.0, float), contingency=(0.0, 50.0, float),
                    gst=(0.0, 40.0, float), cess=(0.0, 10.0, float), cost_index=(0.1, 10.0, float),
                    column_spacing_area=(4.0, 60.0, float), cement_rate=(0.0, 1e6, float), steel_rate=(0.0, 1e7, float),
                    land_area=(0.0, 1e7, float), land_rate=(0.0, 1e9, float), stamp_duty_pct=(0.0, 20.0, float))
CHOICES = dict(finish=("economy", "standard", "premium"))
MAX_QTY, MAX_RATE = 1e9, 1e9


def checked_number(name, v, lo, hi, cast=float):
    try:
        x = cast(float(v))
    except (TypeError, ValueError):
        raise HTTPException(400, f"{name} must be a number")
    if not (lo <= x <= hi) or x != x:
        raise HTTPException(400, f"{name} must be between {lo:g} and {hi:g}")
    return x


DERIVABLE = ("floors", "floor_height", "plinth_height", "foundation_depth", "slab_thickness", "parapet_height", "sill_height", "lintel_height")


# ============================================================================ storage
def db():
    con = sqlite3.connect(DB, check_same_thread=False, timeout=60)
    con.row_factory = sqlite3.Row
    return con


def init_db():
    existed = os.path.exists(DB)
    with db() as con:
        legacy = existed and not con.execute("SELECT 1 FROM sqlite_master WHERE name='rate_schedules'").fetchone()
    if legacy:  # first start after the jurisdiction upgrade: keep a copy of the database as it was
        shutil.copy(DB, DB + f".pre-schedules-{datetime.now():%Y%m%d%H%M%S}.bak")
    with db() as con:
        con.execute("""CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY, name TEXT, client TEXT, location TEXT,
                       created TEXT, updated TEXT, status TEXT, params TEXT, sheets TEXT, items TEXT, overrides TEXT,
                       history TEXT, versions TEXT)""")
        tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if "sor" in tables:
            cols = {r[1] for r in con.execute("PRAGMA table_info(sor)")}
            if "analysis" not in cols:
                con.execute("ALTER TABLE sor ADD COLUMN analysis TEXT")
        cols = {r[1] for r in con.execute("PRAGMA table_info(projects)")}
        if "versions" not in cols:
            con.execute("ALTER TABLE projects ADD COLUMN versions TEXT")
        if "details" not in cols:  # project details (client, building type, site address, description)
            con.execute("ALTER TABLE projects ADD COLUMN details TEXT")
        if "scope" not in cols:  # estimation scope (room / section / floors / property / custom) — NULL = whole property
            con.execute("ALTER TABLE projects ADD COLUMN scope TEXT")
        rates.ensure_schema(con, DSR_CSV, _seed_fallback)
        tn_scope.bootstrap(con)  # TN-only: other schedules inactive (kept for legacy audit); TN edition from snapshot if missing


DSR_CSV = os.path.join(REPO_DATA, "sor", "cpwd_dsr_2023_civil.csv")
LIB = rates.Library()


def _seed_fallback():
    """Built-in sample library, used only when the CPWD DSR file is missing."""
    for c, d, u, r, k in SEED_SOR:
        a = SEED_ANALYSIS.get(c)
        yield (c, d, u, analysis_rate(a) if a else r, k, "analysis" if a else "seed"), (json.dumps(a) if a else None)


def library(sid):
    """(rate items by code, matcher) of one schedule — only that schedule is ever searched."""
    with db() as con:
        return LIB.get(con, sid)


def schedule(sid):
    with db() as con:
        return rates.get_schedule(con, sid)


def load_project(pid):
    with db() as con:
        r = con.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
    if not r:
        raise HTTPException(404, "Project not found")
    p = dict(r)
    for k in ("params", "sheets", "items", "overrides", "history", "versions"):
        p[k] = json.loads(p[k] or ("{}" if k in ("params", "overrides") else "[]"))
    p["loc"] = json.loads(p.get("loc") or "{}")
    p["rate_selection"] = json.loads(p.get("rate_selection") or "{}")
    p["details"] = json.loads(p.get("details") or "null") or dict(client=p.get("client") or "")
    sc = json.loads(p.get("scope") or "null")
    p["scope"] = sc if isinstance(sc, dict) and sc.get("type") else scope_mod.default()
    return p


def save_project(p):
    p["updated"] = datetime.now().isoformat(timespec="seconds")
    with db() as con:
        con.execute("""INSERT OR REPLACE INTO projects(id, name, client, location, created, updated, status, params, sheets, items,
                       overrides, history, versions, loc, rate_schedule_id, rate_selection, details, scope) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (p["id"], p["name"], p.get("client", ""), p.get("location", ""), p["created"], p["updated"],
                     p.get("status", "draft"), json.dumps(p["params"]), json.dumps(p["sheets"], default=_np),
                     json.dumps(p["items"], default=_np), json.dumps(p["overrides"]), json.dumps(p.get("history", [])),
                     json.dumps(p.get("versions", []), default=_np), json.dumps(p.get("loc") or {}),
                     p.get("rate_schedule_id"), json.dumps(p.get("rate_selection") or {}), json.dumps(p.get("details") or {}),
                     json.dumps(p.get("scope") or scope_mod.default())))


def _np(o):
    if isinstance(o, np.integer):
        return int(o)
    if isinstance(o, np.floating):
        return float(o)
    if isinstance(o, np.ndarray):
        return o.tolist()
    if isinstance(o, np.bool_):
        return bool(o)
    raise TypeError(type(o))


def clean(o):
    return json.loads(json.dumps(o, default=_np))


def editable(p):
    if p.get("status") == "confirmed":
        raise HTTPException(409, "This estimate is confirmed and locked. Reopen it to make changes.")


def log(p, text):
    p.setdefault("history", []).append(dict(t=datetime.now().isoformat(timespec="seconds"), text=text))


# ============================================================================ pipeline
def _prefix_ids(sheet, tag):
    for grp in ("rooms", "walls", "openings", "columns"):
        for e in sheet.get(grp) or []:
            e["id"] = f"{tag}-{e['id']}"


SITE_RX = re.compile(r"\bSITE\s*(?:LAYOUT|PLAN)\b|\bMASTER\s*PLAN\b|\bKEY\s*PLAN\b|\bLAYOUT\s+PLAN\s+OF\s+(?:CAMPUS|SITE)\b|\bCAMPUS\s+(?:LAYOUT|PLAN)\b", re.I)


def site_plan_reason(sheet: dict) -> str | None:
    """A site / campus layout (many buildings, roads, 1:500+ scale) must never be costed as a floor plan."""
    blob = " ".join(sheet.get("texts") or []) + " " + str(sheet.get("title") or "")
    den = (sheet.get("notes") or {}).get("scale_den")
    m = re.search(r"\b1\s*[:/]\s*(\d{3,5})\b", blob)
    den = den or (int(m.group(1)) if m else None)
    plan_words = len(re.findall(r"(?:GROUND|FIRST|SECOND|TYPICAL|BASEMENT)\s*FLOOR\s*PLAN", blob, re.I))  # OCR often drops spaces
    if den and den >= 500 and not plan_words:
        return f"printed scale 1:{den} — site / layout scale, not a building floor plan"
    t = SITE_RX.search(blob)
    if t and not plan_words:
        return f"sheet titled '{t.group(0)}' — site / campus layout"
    return None


FLOOR_WORDS = {"GROUND": 0, "FIRST": 1, "SECOND": 2, "THIRD": 3, "FOURTH": 4, "FIFTH": 5}


def written_floor_areas(texts: list) -> dict:
    """Per-floor plinth areas printed on a sheet: 'PLINTH AREA: GROUND FLOOR - 233.83 SQ.M', 'FIRST FLOOR - 202.43 SQ.M'
    (OCR forms such as 'GROUNDFL00R-233.83SQ.M' included). {level: m²}"""
    out = {}
    blob = " ".join(str(t) for t in texts).upper()
    blob = re.sub(r"FL[O0]{2}R", "FLOOR", blob)
    for m in re.finditer(r"(GROUND|FIRST|SECOND|THIRD|FOURTH|FIFTH)\s*FLOOR\s*[-:=]?\s*(\d{2,5}(?:\.\d+)?)\s*SQ\.?\s*M", blob):
        out.setdefault(FLOOR_WORDS[m.group(1)], float(m.group(2)))
    return out


def derived_floors(r: dict) -> list:
    """A sheet that prints the plinth areas of several floors but holds one measured plan (e.g. TNRD community hall III:
    ground and first floor side by side). The measured plan is matched to its floor by area; each other floor is added
    as a DERIVED floor: the measured plan scaled to that floor's written area (lengths × k, areas × k²), scale status
    NEEDS REVIEW — never presented as measured. Returns the extra sheets."""
    if r.get("kind") != "floor_plan" or not r.get("rooms") or r.get("area_only"):
        return []
    fa = written_floor_areas(list(r.get("texts") or []) + [str(r.get("title") or "")])
    plinth = (r.get("metrics") or {}).get("plinth_area") or 0
    if len(fa) < 2 or plinth <= 0:
        return []
    own = min(fa, key=lambda lv: abs(fa[lv] / plinth - 1))
    if abs(fa[own] / plinth - 1) > 0.05:  # the measured plan does not match any written floor area: do not guess
        return []
    name = lambda lv: "Ground floor" if lv == 0 else f"{['First', 'Second', 'Third', 'Fourth', 'Fifth'][lv - 1]} floor"
    r["floor"] = dict(label=name(own), level=own, repeat=1, occurrences=[own])
    if r.get("floor_info"):  # the raw text reading ("FIRST FLOOR PLAN" printed beside it) no longer decides this plan's floor
        r["floor_info"] = dict(r["floor_info"], floor_occurrences=[own], label=name(own))
    r["floor_from_written_area"] = f"matched to the written {name(own).lower()} plinth area {fa[own]:.2f} m²"
    import copy
    import math
    out = []
    for lv, a in sorted(fa.items()):
        if lv == own:
            continue
        k2 = a / plinth
        k = math.sqrt(k2)
        e = copy.deepcopy(r)
        e["sheet"] = f"{r.get('sheet', 'Sheet')} · {name(lv)} (derived)"
        e["floor"] = dict(label=name(lv), level=lv, repeat=1, occurrences=[lv])
        e["derived_floor"] = (f"{name(lv)}: not drawn as a separate measurable view — derived from the {name(own).lower()} plan, "
                              f"scaled to the written {name(lv).lower()} plinth area {a:.2f} m² (×{k2:.3f}); review its rooms and openings")
        m = e["metrics"]
        for key in ("L_ext", "L_int", "envelope_perimeter"):
            if key in m:
                m[key] = round(m[key] * k, 3)
        for key in ("carpet_area",):
            if key in m:
                m[key] = round(m[key] * k2, 3)
        m["plinth_area"] = round(a, 2)
        for rm in e.get("rooms") or []:
            rm["area"] = round((rm.get("area") or 0) * k2, 3)
            if rm.get("perimeter"):
                rm["perimeter"] = round(rm["perimeter"] * k, 3)
        e["scale"] = dict(e.get("scale") or {}, method="derived floor: " + e["derived_floor"], confidence=0.5)
        e.pop("preview_scale", None)
        out.append(e)
    return out


def area_only_plan(sheet: dict) -> float | None:
    """Open structures (threshing floor, cremation shed, shelters) have no enclosed rooms, so the room-based reader finds
    no plan — but their sheets say "PLAN" and print the plinth area. Returns that written area (m²) for a sheet that was
    not read as a measured floor plan, else None."""
    if sheet.get("kind") in ("site_plan", "structural") or (sheet.get("kind") == "floor_plan" and sheet.get("rooms")):
        return None
    texts = list(sheet.get("texts") or []) + [str(sheet.get("title") or "")]
    if not re.search(r"\bPLAN\b", " ".join(texts), re.I):
        return None
    import sheetreader
    pa = sheetreader.plinth_area_m2(texts)
    # the whole estimate would rest on this one number: accept it only when the m² and sq ft figures printed together
    # agree (an OCR fragment such as "PLINTH AREA :2" is not enough)
    return round(pa[0], 2) if pa and pa[1] == "m² and sq ft agree" and 2 <= pa[0] <= 20000 else None


def apply_area_only(r: dict, area_m2: float) -> None:
    """An area-only plan: plinth = the written area (measured); no rooms, walls or openings are invented."""
    r["kind"], r["area_only"] = "floor_plan", True
    r["area_only_reason"] = (f"open / area-only structure: no enclosed rooms were found, the sheet states a plinth area of "
                             f"{area_m2:.2f} m²; only plinth-area work is quantified — walls, roof and frame are not measured")
    r["discarded_rooms"] = (r.get("discarded_rooms") or []) + [dict(id=x.get("id"), area=x.get("area"), reason="area-only sheet")
                                                                for x in r.get("rooms") or []]
    r["rooms"], r["openings"] = [], []
    m = r.setdefault("metrics", {})
    # the platform / floor base of an open structure is its plinth area (carpet = plinth); walls are not measured
    m.update(plinth_area=area_m2, carpet_area=area_m2, envelope_perimeter=0.0, n_doors=0, n_windows=0, L_ext=0.0, L_int=0.0)
    m.setdefault("t_ext", 0.23)
    m.setdefault("t_int", 0.115)
    for k in ("walls", "columns"):
        r.setdefault(k, [])
    sc = r.get("scale") or {}
    r["scale"] = dict(sc, method=f"from the plinth area note ({area_m2:.2f} m²) — area-only open structure", confidence=0.85)


def structural_detail_reason(sheet: dict) -> str | None:
    """Reinforcement / joint detail sheets (bars, stirrups, laps) with only unlabelled 'spaces' are not floor plans:
    their boxes and bar hooks must not become rooms, doors or windows (stress-pack REF_229 / REF_230)."""
    blob = " ".join(sheet.get("texts") or []) + " " + str(sheet.get("title") or "")
    hits = len(re.findall(r"\b(?:BARS?|STIRRUPS?|REINFORCEMENT|LAP(?:PING)?|DEVELOPMENT LENGTH|JOINT DETAILS?|DIA|HOOKS?)\b", blob, re.I))
    named = [x for x in sheet.get("rooms") or [] if not str(x.get("name", "")).startswith("SPACE")]
    if hits >= 6 and not named and not re.search(r"FLOOR\s*PLAN|\bPLAN\b", blob, re.I):
        return f"reinforcement detail sheet ({hits} bar / stirrup / joint terms, no labelled rooms) — not a floor plan"
    return None


def analyze_files(pid: str, paths: list[tuple[str, str]], opts: dict | None = None, start: int = 0) -> list[dict]:
    sheets = []
    os.makedirs(os.path.join(PREVIEWS, pid), exist_ok=True)
    for path, orig in paths:
        t0 = time.time()
        try:
            res = parsers.parse_any(path, STATE["models"], opts or {})
        except Exception as e:
            sheets.append(dict(file=orig, sheet="-", kind="other", error=str(e), rooms=[], walls=[], openings=[], columns=[], metrics={},
                               path=os.path.relpath(path, DATA)))
            continue
        for r in res:
            prev = r.pop("_preview", None)
            fn = f"{uuid.uuid4().hex[:8]}.png"
            if prev is not None and getattr(prev, "size", 0):
                cv2.imwrite(os.path.join(PREVIEWS, pid, fn), prev)
                r["preview"] = f"/api/previews/{pid}/{fn}"
                r["preview_size"] = [int(prev.shape[1]), int(prev.shape[0])]
            r["file"] = orig
            r["path"] = os.path.relpath(path, DATA)
            r["seconds"] = round(time.time() - t0, 2)
            r.setdefault("kind", "floor_plan" if len(r.get("rooms", [])) >= 2 else "other")
            # site / campus / master plans are not building floor plans: no construction quantities from them
            sp = site_plan_reason(r)
            if sp:
                r["kind"], r["site_plan_reason"] = "site_plan", sp
            sd = structural_detail_reason(r) if r["kind"] == "floor_plan" else None
            if sd:
                r["kind"], r["structural_detail_reason"] = "structural", sd
                r["discarded_rooms"] = (r.get("discarded_rooms") or []) + [dict(id=x.get("id"), area=x.get("area"), reason=sd) for x in r.get("rooms") or []]
                r["rooms"], r["openings"] = [], []
            ao = area_only_plan(r)
            if ao:
                apply_area_only(r, ao)
            # floor occurrences stated on the sheet ("18TH, 28TH & 38TH FLOOR", "TYPICAL FLOOR PLAN (5TH TO 12TH)" …)
            fi = floors.parse(r.get("texts") or [], r.get("title"))
            if fi:
                r["floor_info"] = fi
                if fi.get("floor_occurrences") and (not r.get("floor") or len(fi["floor_occurrences"]) > int((r.get("floor") or {}).get("repeat", 1))):
                    r["floor"] = dict(label=fi["label"], level=fi["level"], repeat=fi["repeat"], occurrences=fi["floor_occurrences"])
            drop_annotation_artefacts(r)
            rooms_mod.annotate(r)
            openings.enrich(r)
            r["areas"] = areas.compute(r)
            r.pop("frame", None)
            extra = derived_floors(r)  # other floors whose plinth area the same sheet prints
            _prefix_ids(r, f"S{start + len(sheets) + 1}")
            sheets.append(clean(r))
            for e in extra:
                _prefix_ids(e, f"S{start + len(sheets) + 1}")
                sheets.append(clean(e))
    return sheets


def drop_annotation_artefacts(r: dict):
    """Legend swatches / title-block symbols outside the plan get segmented as small unlabelled 'rooms'. An unlabelled
    space < 2.5 m² whose centre lies outside the extent of the labelled rooms (+0.5 m) is removed from rooms, carpet and
    plinth area, and recorded in r['discarded_rooms'] (never silently)."""
    rooms = r.get("rooms") or []
    named = [x for x in rooms if not str(x.get("name", "")).startswith("SPACE") and x.get("bbox_m")]
    if len(named) < 2:
        return
    x0 = min(x["bbox_m"][0] for x in named) - 0.5
    y0 = min(x["bbox_m"][1] for x in named) - 0.5
    x1 = max(x["bbox_m"][0] + x["bbox_m"][2] for x in named) + 0.5
    y1 = max(x["bbox_m"][1] + x["bbox_m"][3] for x in named) + 0.5
    keep, drop = [], []
    for x in rooms:
        cx, cy = x.get("cx"), x.get("cy")
        if re.fullmatch(r"\s*(?:SECTION|SEC\.?|ELEVATION|DETAIL)\s*[-:]?\s*[A-Z0-9]{1,3}\s*[-–]?\s*[A-Z0-9]{0,3}\s*", str(x.get("name", "")), re.I):
            drop.append(dict(id=x.get("id"), area=x.get("area"), reason=f"view label '{x.get('name')}' is not a room"))
        elif (str(x.get("name", "")).startswith("SPACE") and (x.get("area") or 0) < 2.5 and cx is not None
                and not (x0 <= cx <= x1 and y0 <= cy <= y1)):
            drop.append(dict(id=x.get("id"), area=x.get("area"), reason="small unlabelled space outside the labelled plan "
                                                                       "(legend / title-block artefact)"))
        else:
            keep.append(x)
    if drop:
        r["rooms"] = keep
        r["discarded_rooms"] = drop
        m = r.setdefault("metrics", {})
        a = sum(d["area"] or 0 for d in drop)
        m["n_rooms"] = len(keep)
        if m.get("carpet_area"):
            m["carpet_area"] = round(m["carpet_area"] - a, 2)
        if m.get("plinth_area"):
            m["plinth_area"] = round(m["plinth_area"] - a, 2)


def _rescale_sheet(p, idx, k, note):
    """Re-run one sheet's analysis with its length scale multiplied by k (real / measured) — user calibration and the
    scale guard share this path."""
    s = p["sheets"][idx]
    opts = {}
    if s.get("source") == "pdf":
        opts["scale_den"] = int(round((s["scale"].get("scale_den") or 100) * k))
    elif s.get("source") == "image":
        opts["px_per_m"] = s["scale"]["px_per_m"] / k
    else:
        return None
    path = os.path.join(DATA, s["path"])
    with LOCK:
        new = analyze_files(p["id"], [(path, s["file"])], opts, start=idx)
    page = [n for n in new if n.get("sheet") == s.get("sheet")] or new
    if not page:
        return None
    page[0]["scale"]["method"] = note + " — " + page[0]["scale"]["method"]
    page[0]["floor"] = s.get("floor") or page[0].get("floor")
    for key in ("enabled", "user_label", "floor_confirmed", "opening_confirmed", "kind_user"):
        if key in s:
            page[0][key] = s[key]
    p["sheets"][idx] = page[0]
    return page[0]


def reconcile_area_only(p):
    """Area-only sheets stand in for a plan only when the project has no measured floor plan; next to a real plan they
    would count the same building twice, so they are set aside (kept visible, with the reason)."""
    real = any(s.get("kind") == "floor_plan" and s.get("rooms") and s.get("enabled", True) for s in p["sheets"])
    for s in p["sheets"]:
        if s.get("area_only") and real and s.get("kind") == "floor_plan":
            s["kind"] = "other"
            s["area_only_reason"] = "set aside: the project has a measured floor plan (this sheet only repeats a written plinth area)"


def apply_scale_guard(p):
    """ASSUMED raster scales are checked against the user's built-up area, then against typical room sizes (scale_guard)."""
    import scale_guard
    reconcile_area_only(p)
    for s in p["sheets"]:
        if s.get("kind") == "floor_plan" and s.get("rooms") is not None:
            s["areas"] = areas.compute(s)
    plans = [i for i, s in enumerate(p["sheets"]) if s.get("kind") == "floor_plan" and s.get("enabled", True)]
    try:
        user_m2 = float(((p.get("scope") or {}).get("config") or {}).get("built_up_area_m2") or 0) or None
    except (TypeError, ValueError):
        user_m2 = None
    for i in plans:
        s = p["sheets"][i]
        if not scale_guard.needs_guard(s):
            continue
        f = scale_guard.user_area_factor(s, user_m2, len(plans))
        if not f:
            # typical room sizes NEVER re-measure the drawing (that would fabricate walls, concrete, steel): they only
            # flag the scale and give a separately labelled preliminary area for a budget range
            pr = scale_guard.prior_factor(s)
            plinth = (s.get("metrics") or {}).get("plinth_area") or 0
            s["scale_guard"] = dict(applied=False) if not pr else dict(
                applied=False, suspect=True, reason=pr[1], suggested_factor=round(pr[0], 3),
                implied_plinth_m2=round(plinth * pr[0] ** 2, 1))
            continue
        k, why = f
        new = _rescale_sheet(p, i, k, ("calibrated from " if why.startswith("user") else "") + why)
        if new is not None:
            new["scale_guard"] = dict(applied=True, factor=round(k, 3), reason=why)
            new["areas"] = areas.compute(new)
            log(p, f"Scale corrected on {new['file']}: {why}")


def refresh_derived(p):
    """Parameters read from the drawings (heights, floors …) unless the user has set them."""
    plans = [s for s in p["sheets"] if s.get("enabled", True)]
    for s in p["sheets"]:  # area definitions + scale status follow user calibration / scale confirmation
        if s.get("kind") == "floor_plan" and s.get("rooms") is not None:
            s["areas"] = areas.compute(s)
    vals, src = takeoff.derive_params(plans)
    srcs = p["params"].setdefault("_src", {})
    for k in DERIVABLE:
        if srcs.get(k) == "user":
            continue
        if k in vals:
            p["params"][k] = vals[k]
            srcs[k] = src[k]
        else:
            p["params"][k] = takeoff.DEFAULT_PARAMS[k]
            srcs[k] = "default"


_MAPS = {}


def curated_map(code_system):
    """Hand-checked take-off -> official item map of a schedule numbering system (data/sor/maps/<system>.json)."""
    if code_system not in _MAPS:
        f = os.path.join(REPO_DATA, "sor", "maps", f"{code_system}.json")
        _MAPS[code_system] = {k: v for k, v in json.load(open(f, encoding="utf-8")).items() if not k.startswith("_")} if os.path.exists(f) else {}
    return _MAPS[code_system]


def map_items(p, items):
    """Match take-off items against the project's rate schedule only (never another jurisdiction's).
    - CPWD-numbered schedules: expert table + NLP;
    - schedules that publish finished item rates: NLP with a rate-plausibility check, always for review;
    - schedules that publish BASIC rates (materials, labour): rate analysis from those official basic rates."""
    import analysis
    sid = p.get("rate_schedule_id")
    by, matcher = library(sid)
    s = schedule(sid) if sid else None
    cs = (s or {}).get("code_system") or ""
    if s and cs == "TN-SOR":
        with db() as con:
            tn_pricing.price_items(con, items, by, s)
            import mapping_review as MR
            approved_maps = MR.approved(con, s["id"])
            for it in items:  # engineer-approved mappings win over automatic suggestions
                ap = approved_maps.get((it["key"], it["unit"]))
                if ap and str(ap["selected_code"]).startswith("SD:"):
                    src_no = ((it.get("derived") or {}).get("source") or {}).get("data_no")
                    if src_no and f"SD:{src_no}" == ap["selected_code"]:
                        it.update(map_method="engineer-approved", status="confirmed",
                                  map_check=f"Standard Data {src_no} analysis approved by {ap['reviewer']} on {ap['at'][:10]}")
                    else:
                        it.update(status="review", map_check=f"engineer approved {ap['selected_code']} but no complete analysis of that data "
                                                              f"number is available — NEEDS RATE ANALYSIS review")
                    continue
                if ap and ap["selected_code"] in by:
                    r_ = by[ap["selected_code"]]
                    it.update(sor_code=r_["code"], map_method="engineer-approved", status="confirmed", rate_gap=None, map_score=1.0,
                              candidates=[dict(code=r_["code"], description=r_["description"][:160], unit=r_["unit"], rate=r_["rate"], score=1.0)],
                              map_check=f"engineer-approved by {ap['reviewer']} on {ap['at'][:10]}" + (f" — unit note: {ap['unit_justification']}" if ap.get("unit_justification") else ""))
                    it.pop("derived", None)
            learning.apply_learned(con, items, by)  # what engineers confirmed on earlier estimates
            tn_pricing.apply_estimates(items)  # last resort: editable average rate for lines still without a PWD rate
        return
    if s and cs in analysis.BASIC_RATE_SYSTEMS:
        for it in items:
            zone = (p.get("params") or {}).get("labour_zone") or "Z1"
            d = analysis.derive(it["key"], cs, by, zone, p.get("params")) if not it.get("custom") else None
            it.update(candidates=[], sor_code=None, map_score=0, nlp_agrees_with_rules=False)
            if d and d["ok"] and d["unit"] == it["unit"]:
                it.update(derived=d, map_method="analysis", status="auto", map_score=1.0)
            else:
                it.pop("derived", None)
                it.update(map_method="none", status="unmapped",
                          map_check=("official basic rate missing in this schedule: " + ", ".join(d["missing"])
                                     + (" — enter the current notified rate under 'Cement / steel rate'" if set(d["missing"]) & {"cement", "steel"} else "")) if d and not d["ok"]
                          else "no official item rate or analysis for this item in this schedule — select or add a rate")
        return
    expert = bool(s and cs == "CPWD-DSR")
    ref = None if expert else takeoff.reference_rates(library(rates.CPWD_ID)[0])
    cmap = curated_map(cs) if s else {}
    rest = []
    for it in items:
        e = cmap.get(it["key"]) if not it.get("custom") else None
        for f in ("sor_factor", "included", "derived"):
            it.pop(f, None)
        if e and e.get("unpriced"):
            it.update(candidates=[], sor_code=None, map_score=0, map_method="none", status="unmapped",
                      map_check=e["unpriced"], nlp_agrees_with_rules=False)
        elif e and e.get("included"):
            it.update(candidates=[], sor_code=None, map_score=1.0, map_method="included", status="auto",
                      included=e.get("note") or f"included in {e['included']}", nlp_agrees_with_rules=False)
        elif e and e.get("code") in by:
            r = by[e["code"]]
            it.update(candidates=[dict(code=r["code"], description=r["description"][:140], unit=r["unit"], rate=r["rate"], score=1.0)],
                      sor_code=r["code"], map_score=1.0, map_method="table", status="auto", nlp_agrees_with_rules=False,
                      map_check=e.get("note", ""))
            if e.get("factor"):
                it["sor_factor"] = float(e["factor"])
        else:
            rest.append(it)
    takeoff.map_rates(rest, matcher, by, expert=expert, ref=ref)


def rebuild_items(p):
    usable = [s for s in p["sheets"] if s.get("enabled", True)]
    eff = scope_mod.effective_params(p["params"], p.get("scope") or scope_mod.default())  # room scope: never × floors
    items = takeoff.compute_quantities(usable, eff)  # quantities: drawings + parameters only, never location
    # a bar bending schedule, when supplied, is THE reinforcement quantity for its member type (replaces, never adds to,
    # the plan-schedule / kg-per-m³ figure — no double counting)
    import structural as structural_mod
    sq = structural_mod.compute(usable, {**takeoff.DEFAULT_PARAMS, **eff}, (p["params"] or {}).get("_src"))
    for it in items:
        mt = it["key"].replace("steel_", "") if it["key"].startswith("steel_") else None  # footing / column / beam / slab
        if mt and mt in sq["steel_by_member_type"]:
            b_ = sq["steel_by_member_type"][mt]
            it.update(qty=b_["kg"], formula=f"Σ BBS rows ({b_['rows']}) L × n × Ø²/162", basis="bar bending schedule (bar by bar)",
                      sources=[r["mark"] or "" for r in sq["bbs_rows"]][:50])
    custom = [i for i in p.get("items", []) if i.get("custom")]
    map_items(p, items + custom)
    p["items"] = items + custom


def effective_overrides(p):
    """Quantity edits always apply; a chosen PWD item / rate applies only within the schedule it was chosen in."""
    sid = p.get("rate_schedule_id")
    out = {}
    for k, o in (p.get("overrides") or {}).items():
        o = dict(o)
        if o.get("schedule_id") and o["schedule_id"] != sid:
            for f in ("sor_code", "rate", "status"):
                o.pop(f, None)
        out[k] = o
    return out


def price_project(p):
    by, _ = library(p.get("rate_schedule_id"))
    sc = p.get("scope") or scope_mod.default()
    ctx = scope_mod.context(p["sheets"])
    return takeoff.price(p["items"], by, effective_overrides(p), p["params"], exclude=lambda r: scope_mod.exclusion(r, sc, ctx))


def set_schedule(p, sid, mode, reason="", by=None, resolved=None):
    """Assign the project's rate schedule (or none) and record the decision; quantities are untouched."""
    old = p.get("rate_schedule_id")
    s = schedule(sid) if sid else None
    sel = p.get("rate_selection") or {}
    hist = sel.get("history", [])
    if old != sid or not sel:
        hist = hist + [dict(at=datetime.now().isoformat(timespec="seconds"), from_schedule=old, to_schedule=sid, mode=mode,
                            reason=reason, by=by or "")]
    p["rate_schedule_id"] = sid
    p["rate_selection"] = dict(mode=mode, status="selected" if sid else (resolved or {}).get("status", "unavailable"),
                               reason=reason or (resolved or {}).get("reason", ""), by=by or "",
                               at=datetime.now().isoformat(timespec="seconds"), history=hist[-30:])
    if old != sid:
        log(p, f"Rate schedule: {rates.label(schedule(old)) if old else 'none'} → {rates.label(s) if s else 'none'} ({mode})"
               + (f" — {reason}" if reason else ""))
    return old != sid


def schedule_view(p):
    s = schedule(p.get("rate_schedule_id"))
    sel = p.get("rate_selection") or {}
    v = dict(location=p.get("loc") or {}, location_text=rates.loc_text(p.get("loc")), schedule=rates.public(s),
             selection={k: sel.get(k) for k in ("mode", "status", "reason", "by", "at")}, history=sel.get("history", []))
    if not s:
        v["message"] = sel.get("reason") or "No official rate schedule is configured for this jurisdiction."
    return v


def full_project(p):
    from reports import trace_text
    est = price_project(p)
    rs = schedule_view(p)
    s = rs["schedule"]
    for it in est["items"]:
        it["trace"] = trace_text(it, p["sheets"])
        priced = s and it.get("rate") is not None
        it["rate_schedule"] = s["label"] if priced else None
        it["rate_schedule_id"] = s["id"] if priced else None
        it["rate_authority"] = s.get("authority") if priced else None
        it["rate_year"] = " / ".join(x for x in (s.get("year"), s.get("financial_year"), ("rev. " + s["revision"]) if s.get("revision") else None) if x) if priced else None
        it["rate_source"] = (s.get("source_ref") or s.get("source")) if priced else None
        it["rate_kind"] = ("estimated average rate (editable)" if it.get("map_method") == "estimated" else
                           "TN Standard Data analysis" if it.get("map_method") == "analysis" else "included in another item" if it.get("map_method") == "included"
                           else "user rate" if it.get("map_method") == "user" else "TN SOR item rate") if priced else None
        prov = _rate_provenance(it, s) if priced else None
        it["rate_provenance"] = prov
        if prov:
            it["rate_page"] = prov.get("page")
            it["rate_effective_from"] = prov.get("effective_from")
            it["rate_sha256"] = prov.get("sha256")
        b = str(it.get("basis") or "drawing")
        it["basis_class"] = ("assumed" if b.startswith("assumed") else "partly assumed" if "assumed" in b else "user input" if "user" in b
                             else "explicit" if re.search(r"schedule|bar by bar|written|label", b) else "derived")
        it["rate_gap"] = None if it.get("rate") is not None else (it.get("rate_gap") or "NEEDS RATE")
    plans = [s for s in p["sheets"] if s.get("kind") == "floor_plan" and s.get("enabled", True)]
    bua = 0.0
    if plans:
        plans = takeoff.sort_plans(plans, p["sheets"])
        mult = takeoff.plan_multipliers(plans, max(1, int(p["params"]["floors"])))
        bua = sum(s["metrics"]["plinth_area"] * m for s, m in zip(plans, mult))  # duplicates count 0
    # area for a PRELIMINARY BUDGET RANGE only (never used for quantities): suspect scales replaced by the room-size estimate
    budget_area = None
    if plans and any((s.get("scale_guard") or {}).get("suspect") for s in plans):
        budget_area = round(sum(((s.get("scale_guard") or {}).get("implied_plinth_m2") or s["metrics"]["plinth_area"]) * m
                                for s, m in zip(plans, mult)), 1)
    summary = dict(budget_area_m2=budget_area,
        plinth_area=round(plans[0]["metrics"]["plinth_area"], 2) if plans else 0, built_up_area=round(bua, 2),
        carpet_area=round(sum(s["metrics"]["carpet_area"] for s in plans), 2),
        rooms=sum(len(s["rooms"]) for s in plans), doors=sum(s["metrics"]["n_doors"] for s in plans),
        windows=sum(s["metrics"]["n_windows"] for s in plans), floor_plans=len(plans),
        structural=sum(1 for s in p["sheets"] if s.get("kind") == "structural" and s.get("enabled", True)),
        sections=sum(1 for s in p["sheets"] if s.get("kind") in ("section", "elevation") and s.get("enabled", True)),
        cost_per_sqft=round(est["total"] / (bua * 10.764)) if bua else None,
        cost_range=__import__("scale_guard").cost_range(est["total"], [s_ for s_ in p["sheets"] if s_.get("enabled", True)]),
    )
    # two figures: (1) construction cost from the rate schedule - no land; (2) the same plus land at a rate the
    # user enters (official guideline value or their own estimate). No market data is fetched.
    pp = {**takeoff.DEFAULT_PARAMS, **p["params"]}
    la, lr = float(pp.get("land_area") or 0), float(pp.get("land_rate") or 0)
    land_cost = la * lr
    duty = land_cost * float(pp.get("stamp_duty_pct") or 0) / 100
    land = dict(area_sqm=la, rate_per_sqm=lr, basis=pp.get("land_rate_basis") or "", land_cost=round(land_cost, 2),
                stamp_duty_registration=round(duty, 2), provided=la > 0 and lr > 0,
                construction_cost=est["total"], total_with_land=round(est["total"] + land_cost + duty, 2))
    others = takeoff.other_elements([s for s in p["sheets"] if s.get("enabled", True)])
    structural = any(s_.get("kind") == "structural" or s_.get("structural") for s_ in p["sheets"] if s_.get("enabled", True))
    stats = dict(items=len(est["items"]), excluded=len(est.get("excluded", [])), other_elements=len(others),
                 structural_input="present" if structural else "STRUCTURAL INPUT MISSING — footing, column, beam sizes and steel are ASSUMED",
                 by_basis={k: sum(1 for i in est["items"] if i.get("basis_class") == k) for k in ("explicit", "derived", "partly assumed", "assumed", "user input")},
                 floors_info=[s_.get("floor_info") for s_ in p["sheets"] if s_.get("floor_info")],
                 units=[u for s_ in p["sheets"] for u in (s_.get("units") or [])],
                 review=sum(1 for i in est["items"] if i["status"] in ("review", "unmapped")),
                 validated=sum(1 for i in est["items"] if i["status"] in ("confirmed", "edited")),
                 assumed=sum(1 for i in est["items"] if str(i.get("basis", "")).startswith("assumed")))
    import structural as structural_mod
    structural_q = structural_mod.compute(p["sheets"], {**takeoff.DEFAULT_PARAMS, **p["params"]}, (p["params"] or {}).get("_src"))
    import walls as walls_mod
    for s_ in p["sheets"]:
        if s_.get("kind") == "floor_plan" and s_.get("enabled", True):
            s_["wall_quantities"] = walls_mod.compute(s_, {**takeoff.DEFAULT_PARAMS, **p["params"]}, (p["params"] or {}).get("_src"))
    for it in est["items"]:
        it["review_level"] = quality.review_level(it)
    assumptions = quality.assumptions(p, est)
    checks = quality.checks(p, est)
    stats["completeness"] = quality.completeness(est)
    import confidence as conf_mod
    worst_scale = next((st_ for st_ in ("UNRESOLVED", "NEEDS REVIEW") for s_ in p["sheets"]
                        if s_.get("kind") == "floor_plan" and ((s_.get("areas") or {}).get("calibration") or {}).get("status") == st_), None)
    ic = [conf_mod.item(it, worst_scale) for it in est["items"]]
    for it, c_ in zip(est["items"], ic):
        it["confidence_label"], it["confidence_reasons"] = c_["label"], c_["reasons"]
    stats["estimate_confidence"] = conf_mod.estimate(est, stats["completeness"], structural_q.get("status"), ic)
    scale_ok = not any(s_.get("kind") == "floor_plan" and s_.get("enabled", True)
                       and ((s_.get("areas") or {}).get("calibration") or {}).get("status") != "MEASURED" for s_ in p["sheets"])
    for it in est["items"]:
        it["qty_status"] = quality.qty_status(it, scale_ok)
        it["rate_status"] = quality.rate_status(it)
    stats["status_summary"] = quality.status_summary(p, est, others, scale_ok)
    stats["scope_context"] = scope_mod.context(p["sheets"])
    stats["checks"] = {lv: sum(1 for c in checks if c["level"] == lv) for lv in ("ERROR", "CONFLICT", "WARNING", "INFO")}
    sc_ = p.get("scope") or scope_mod.default()
    return clean(dict(id=p["id"], name=p["name"], client=p.get("client", ""), location=p.get("location", ""), details=p.get("details") or {},
                      scope=sc_, scope_text=scope_mod.describe(sc_),
                      assumptions=assumptions, checks=checks, structural=structural_q,
                      loc=p.get("loc") or {}, rate_source=rs,
                      created=p["created"], updated=p["updated"], status=p.get("status", "draft"), params=p["params"],
                      sheets=p["sheets"], estimate=est, summary=summary, stats=stats, land=land, other_elements=others,
                      versions=[{k: v for k, v in x.items() if k != "items"} for x in p.get("versions", [])],
                      history=p.get("history", [])[-40:]))


def _rate_provenance(it, s):
    d = it.get("derived")
    if it.get("map_method") == "estimated" and d:
        return dict(kind="Estimated average rate", document="estimated average rates", basis=d.get("basis"), verification="estimated",
                    page=None, sha256=None)
    if d and d.get("source"):
        src = d["source"]
        return dict(kind="TN Standard Data analysis", document=src.get("document"), data_no=src.get("data_no"), page=src.get("page"),
                    sha256=src.get("sha256"), url=src.get("url"), effective_from=src.get("effective_from"), verification=src.get("verification"),
                    basic_rates=f"{s.get('label')} (SOR basic rates)", components=d.get("components"))
    code = it.get("sor_code")
    if not code:
        return None
    with db() as con:
        r = con.execute("SELECT source FROM rate_items WHERE schedule_id=? AND code=?", (s["id"], code)).fetchone()
    try:
        src = json.loads(r["source"]) if r and r["source"] and r["source"].startswith("{") else {}
    except ValueError:
        src = {}
    return dict(kind="TN SOR item", code=code, document=src.get("source_title"), edition=src.get("edition"), page=src.get("page"),
                sha256=src.get("base_sha256"), url=src.get("source_url"), effective_from=src.get("effective_from"),
                amended=src.get("amended"), amendment_sha256=src.get("amendment_sha256") if src.get("amended") else None,
                verification=src.get("verification_status"))


# ============================================================================ Tamil Nadu PWD data (admin)
@app.get("/api/tnpwd/status")
def tnpwd_status():
    from tnpwd_ingestion import refresh_service as R
    from tnpwd_ingestion.store import connect
    st = R.status(connect(DB))
    with db() as c2:
        s = tn_scope.active_tn(c2)
        st["active"] = rates.public(rates._with_count(c2, s)) if s else None
        st["previous"] = [rates.public(dict(r)) for r in c2.execute("SELECT * FROM rate_schedules WHERE region=? AND status='inactive'", (tn_scope.STATE,))]
    return clean(st)


@app.get("/api/tnpwd/search")
def tnpwd_search(q: str, kind: str = ""):
    from tnpwd_ingestion import retrieval_index
    from tnpwd_ingestion.store import connect
    return retrieval_index.search(connect(DB), q, kind or None)


@app.post("/api/tnpwd/{action}")
def tnpwd_action(action: str, body: dict | None = None):
    """check | download | extract | build | activate | refresh | compare — public documents only."""
    from tnpwd_ingestion import refresh_service as R
    from tnpwd_ingestion.store import connect
    con = connect(DB)
    body = body or {}
    if action == "compare":
        out = R.compare(con, body["a"], body["b"])
    elif action == "build":
        out = {k: v for k, v in R.build(con).items() if k != "_items"}
    elif action in ("check", "download", "extract", "activate", "refresh"):
        out = getattr(R, action)(con)
    else:
        raise HTTPException(404, "unknown action")
    LIB.invalidate()
    return clean(out)


# ============================================================================ annotation review (TN dataset)
@app.get("/api/annotations/coverage")
def annotation_coverage():
    import annotations as A
    return clean(A.coverage())


@app.get("/api/annotations/pages")
def annotation_pages(state: str = "", split: str = ""):
    import annotations as A
    pages = A._meta("pages.json", {})
    projects = A._meta("projects.json", {})
    out = []
    for pid, pg in sorted(pages.items()):
        st = A.state(pid)
        sp = projects.get(pg.get("project_id"), {}).get("split")
        if (state and st != state) or (split and sp != split):
            continue
        out.append(dict(page_id=pid, project_id=pg.get("project_id"), file=pg["file"], page=pg["page"], state=st, split=sp,
                        drawing_type=pg["drawing_type"], training_eligible=pg.get("training_eligible"), license_status=pg["license_status"]))
    return out


@app.get("/api/annotations/{page_id}")
def annotation_get(page_id: str):
    import annotations as A
    pages = A._meta("pages.json", {})
    if page_id not in pages:
        raise HTTPException(404, "unknown page")
    pp = os.path.join(A.DS, "pseudo_labels", page_id + ".json")
    return clean(dict(page=pages[page_id], state=A.state(page_id), record=A.load(page_id),
                      pseudo=json.load(open(pp)) if os.path.exists(pp) else None, classes=A.CLASSES,
                      image=f"/api/annotations/{page_id}/image"))


@app.get("/api/annotations/{page_id}/image")
def annotation_image(page_id: str):
    import annotations as A
    f = os.path.join(A.DS, "processed", os.path.basename(page_id) + ".png")
    if not os.path.exists(f):
        raise HTTPException(404, "image not found")
    return FileResponse(f)


@app.put("/api/annotations/{page_id}")
def annotation_save(page_id: str, body: dict):
    import annotations as A
    try:
        return clean(A.save_draft(page_id, body.get("objects") or [], (body.get("annotator") or "").strip(), body.get("notes", "")))
    except A.AnnotationError as e:
        raise HTTPException(422, str(e))


@app.post("/api/annotations/{page_id}/from-pseudo")
def annotation_from_pseudo(page_id: str, body: dict):
    import annotations as A
    try:
        return clean(A.start_from_pseudo(page_id, (body.get("annotator") or "").strip()))
    except A.AnnotationError as e:
        raise HTTPException(422, str(e))


@app.post("/api/annotations/{page_id}/review")
def annotation_review(page_id: str, body: dict):
    import annotations as A
    try:
        return clean(A.review(page_id, (body.get("reviewer") or "").strip(), body.get("decision"), body.get("note", "")))
    except A.AnnotationError as e:
        raise HTTPException(422, str(e))


@app.get("/api/annotations-export/coco")
def annotation_export_coco(training_only: bool = True):
    import annotations as A
    return A.export_coco(A.verified_records(training_only))


# ============================================================================ PWD mapping review (engineers)
@app.get("/api/projects/{pid}/mapping-review")
def mapping_review_queue(pid: str):
    import mapping_review as MR
    p = load_project(pid)
    sugg = MR.suggestions(price_project(p)["items"])
    with db() as con:
        ev = MR.evaluate(con, p.get("rate_schedule_id"), sugg)
        ap = MR.approved(con, p.get("rate_schedule_id"))
    for x in sugg:
        a = ap.get((x["key"], x["unit"]))
        x["approved"] = dict(code=a["selected_code"], reviewer=a["reviewer"], at=a["at"]) if a else None
    return clean(dict(queue=sugg, evaluation=ev))


@app.post("/api/projects/{pid}/mapping-review")
def mapping_review_decide(pid: str, body: dict):
    import mapping_review as MR
    p = load_project(pid)
    editable(p)  # a confirmed estimate must be reopened before its mappings / rates change
    by, _ = library(p.get("rate_schedule_id"))
    code = body.get("code")
    code_unit = by[code]["unit"] if code in by else body.get("code_unit")
    if code and not code.startswith("SD:") and code not in by:
        raise HTTPException(422, f"{code} is not an item of the active TN schedule")
    try:
        with db() as con:
            MR.approve(con, body.get("key"), body.get("unit"), code, (body.get("reviewer") or "").strip(), body.get("decision"),
                       p.get("rate_schedule_id"), code_unit, body.get("unit_justification", ""), body.get("note", ""))
    except ValueError as e:
        raise HTTPException(422, str(e))
    rebuild_items(p)
    save_project(p)
    return mapping_review_queue(pid)


# ============================================================================ learning + reference BOQs
@app.get("/api/reference-boq/template")
def reference_boq_template():
    return FileResponse(os.path.join(REPO_DATA, "templates", "reference_boq_template.csv"), filename="reference_boq_template.csv")


@app.post("/api/projects/{pid}/reference-boq")
def upload_reference_boq(pid: str, file: UploadFile = File(...), approval_status: str = Form(""), approval_authority: str = Form(""),
                         approval_date: str = Form(""), source_reference: str = Form(""), drawing_revision: str = Form(""),
                         boq_revision: str = Form(""), currency: str = Form(""), scope: str = Form(""), inclusions: str = Form(""),
                         exclusions: str = Form(""), taxes_included: str = Form(""), source: str = Form("")):
    """Reference BOQ for this project. See boq.py and data/templates/reference_boq_template.csv for the required fields.
    The original file is stored unchanged with its SHA-256; 'approved' requires authority, date and source reference."""
    import boq
    load_project(pid)
    data = file.file.read()
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(413, "BOQ file larger than 20 MB")
    meta = dict(approval_status=approval_status.strip().lower(), approval_authority=approval_authority, approval_date=approval_date,
                source_reference=source_reference or source, drawing_revision=drawing_revision, boq_revision=boq_revision, currency=currency,
                scope=scope, inclusions=inclusions, exclusions=exclusions, taxes_included=taxes_included)
    with db() as con:
        r = boq.upload(con, pid, data, file.filename, meta, os.path.join(DATA, "reference_boqs", pid))
    if not r["stored"]:
        raise HTTPException(422, dict(message="BOQ not accepted — see validation", **r["validation"]))
    return compare_reference_boq(pid) | dict(upload=r)


@app.get("/api/projects/{pid}/reference-boq/compare")
def compare_reference_boq(pid: str):
    import boq
    import mapping_review as MR
    p = load_project(pid)
    est = price_project(p)
    with db() as con:
        ref = boq.latest(con, pid)
        if ref is None:
            raise HTTPException(404, "No reference BOQ uploaded for this project")
        ap = {k: v["selected_code"] for (k, u), v in MR.approved(con, p.get("rate_schedule_id")).items()}
        cmp = boq.compare(ref["rows"], est["items"], ap)
        m = boq.metrics(cmp, est, ref["rows"], ref["meta"])
        learning.ensure(con)
        con.execute("INSERT INTO evaluation_runs(at, project_id, kind, metrics) VALUES(?,?,?,?)",
                    (datetime.now().isoformat(timespec="seconds"), pid, "reference_boq_" + ref["approval_status"], json.dumps(m, default=str)))
    return clean(dict(metrics=m, comparison=cmp, reference=dict(filename=ref["filename"], sha256=ref["sha256"], approval_status=ref["approval_status"],
                                                                meta=ref["meta"], uploaded_at=ref["uploaded_at"])))


@app.get("/api/projects/{pid}/market-check")
def market_check_route(pid: str, refresh: bool = False):
    """Live web reality check of the estimate's ₹/sq ft for this building type (advisory; nothing is changed)."""
    import market_check
    q = full_project(load_project(pid))
    d = q.get("details") or {}
    bt = d.get("building_type") or (q.get("scope") or {}).get("label") or "building"
    loc = q.get("loc") or {}
    place = f"{loc.get('district')} Tamil Nadu" if loc.get("district") else "Tamil Nadu"
    bm = None if refresh else market_check.cached(bt, place)
    if bm is None:  # read the web in the background; the page polls
        market_check.start(bt, place, refresh)
        return dict(pending=True, building_type=bt, place=place)
    S = q.get("summary") or {}
    est = S.get("cost_per_sqft")
    area_basis = "budget area from typical room sizes (drawing scale looks wrong)" if S.get("budget_area_m2") else "measured built-up area"
    sqft = (S.get("budget_area_m2") or S.get("built_up_area") or 0) * 10.7639
    market_total = dict(low=round(bm["low"] * sqft), median=round(bm["median"] * sqft), high=round(bm["high"] * sqft),
                        built_up_sqft=round(sqft), area_basis=area_basis,
                        label="PRELIMINARY BUDGET RANGE (area × market rate) — not a drawing-based BOQ") if bm.get("ok") and sqft else None
    return clean(dict(benchmark=bm, estimate_per_sqft=est, verdict=market_check.verdict(est, bm), market_total=market_total))


def _summary_facts(q: dict) -> dict:
    E, S = q.get("estimate") or {}, q.get("summary") or {}
    ss = (q.get("stats") or {}).get("status_summary") or {}
    cats = sorted(E.get("categories") or [], key=lambda c: -(c.get("amount") or 0))[:4]
    cal = sorted({((s_.get("areas") or {}).get("calibration") or {}).get("status") or "—" for s_ in q.get("sheets", []) if s_.get("kind") == "floor_plan"})
    return dict(project=q.get("name"), building_type=(q.get("details") or {}).get("building_type"), scope=q.get("scope_text"),
                total_inr=round(E.get("total") or 0), likely_range_inr=S.get("cost_range"), built_up_sqft=round((S.get("built_up_area") or 0) * 10.7639),
                cost_per_sqft_inr=S.get("cost_per_sqft"), floors=(q.get("params") or {}).get("floors"),
                top_categories=[dict(name=c["name"], amount_inr=round(c["amount"])) for c in cats], estimate_status=ss.get("status"),
                status_reasons=(ss.get("reasons") or [])[:4], drawing_scale_status=cal,
                warnings=sorted({c["code"] for c in q.get("checks", []) if c["level"] in ("ERROR", "CONFLICT", "WARNING")})[:6],
                rates_from=(q.get("rate_source") or {}).get("schedule", {}).get("label") if (q.get("rate_source") or {}).get("schedule") else None)


@app.get("/api/projects/{pid}/ai-summary")
def ai_summary_route(pid: str):
    """Plain-language explanation of the estimate by the local Meta Llama model (numbers checked against the facts)."""
    import local_llm
    q = full_project(load_project(pid))
    if not local_llm.available():
        return dict(available=False, message=f"Local model not running — start Ollama with {local_llm.model()} to get written summaries.")
    facts = _summary_facts(q)
    try:
        mc = market_check_route(pid)  # cached result only; never waits for a web read
        if (mc.get("benchmark") or {}).get("ok"):
            facts["market_median_per_sqft_inr"] = mc["benchmark"]["median"]
            facts["market_check"] = mc["verdict"].get("text")
    except Exception:
        pass
    r = local_llm.summarise(facts)
    return dict(available=True, **(r or dict(text=None, rejected=True, reason="no reply from the local model")))


@app.get("/api/rate-lookup")
def rate_lookup_route(description: str, unit: str, place: str = "Chennai"):
    """Web price suggestion for an item without a PWD rate (never written to the rate database)."""
    import rate_lookup
    return rate_lookup.lookup(description, unit, place)


@app.get("/api/learning")
def learning_state():
    with db() as con:
        learning.ensure(con)
        return dict(learned=[dict(x) for x in con.execute("SELECT * FROM learned_mappings ORDER BY n DESC, last_at DESC")],
                    evaluation_runs=learning.accuracy_history(con))


def _tn_auto_refresh():
    """Weekly: check the TN PWD sources, download / extract / validate new documents. A validated new edition is
    activated only when TN_AUTO_ACTIVATE=1; otherwise it waits for an administrator (Tamil Nadu PWD data page)."""
    import time as _t
    from tnpwd_ingestion import refresh_service as R
    from tnpwd_ingestion.store import connect
    while True:
        try:
            con = connect(DB)
            last = con.execute("SELECT MAX(at) FROM tn_events WHERE kind='check'").fetchone()[0]
            if not last or (datetime.now() - datetime.fromisoformat(last)).days >= 7:
                new = [d for d in R.check(con) if d["status"] == "NEW"]
                if new:
                    R.download(con)
                    R.extract(con)
                    cand = R.build(con)
                    if cand.get("ok") and os.environ.get("TN_AUTO_ACTIVATE") == "1":
                        R.activate(con, cand)
                        LIB.invalidate()
        except Exception as e:  # network down etc. — never affects estimates
            print("[tn-refresh]", e, flush=True)
        _t.sleep(6 * 3600)


# ============================================================================ startup
@app.on_event("startup")
def _startup():
    init_db()
    samples.ensure_samples(SAMPLES)
    if not os.path.exists(os.path.join(SAMPLES, "drawing_set.pdf")):
        samples.render_set_pdf(os.path.join(SAMPLES, "drawing_set.pdf"))
    STATE["models"] = Models(os.path.join(REPO_DATA, "models"))
    LIB.invalidate()
    if os.environ.get("TN_WEB_REFRESH", "1") == "1" and not os.environ.get("EST_DATA_DIR"):
        threading.Thread(target=_tn_auto_refresh, daemon=True).start()
    print("[server] ready -> http://127.0.0.1:8000", flush=True)


# ============================================================================ API: meta & files
@app.get("/api/health")
def health():
    with db() as con:
        n = con.execute("SELECT COUNT(*) FROM rate_schedules WHERE status='active'").fetchone()[0]
    return dict(ok=True, active_rate_schedules=n)


@app.get("/api/samples")
def list_samples():
    return [
        dict(name="drawing_set", label="Complete drawing set · G+1", desc="Ground & first floor plans, section, foundation plan, column & beam schedules",
             files=["drawing_set.pdf"], thumb="villa_g1.png"),
        dict(name="villa_g1", label="Villa floor plan", desc="Single floor plan · 13.2 × 10.8 m", files=["villa_g1.pdf", "villa_g1.dxf", "villa_g1.png"], thumb="villa_g1.png"),
        dict(name="bungalow", label="Bungalow floor plan", desc="Single floor plan · 14.6 × 11.4 m", files=["bungalow.pdf", "bungalow.dxf", "bungalow.png"], thumb="bungalow.png"),
    ]


@app.get("/api/samples/{fn}")
def get_sample(fn: str):
    path = os.path.join(SAMPLES, os.path.basename(fn))
    if not os.path.exists(path):
        raise HTTPException(404)
    return FileResponse(path, filename=os.path.basename(fn))


@app.get("/api/previews/{pid}/{fn}")
def preview(pid: str, fn: str):
    path = os.path.join(PREVIEWS, os.path.basename(pid), os.path.basename(fn))
    if not os.path.exists(path):
        raise HTTPException(404)
    return FileResponse(path, media_type="image/png")


# ============================================================================ API: projects
@app.get("/api/projects")
def list_projects():
    with db() as con:
        rows = con.execute("SELECT id FROM projects ORDER BY updated DESC").fetchall()
    out = []
    for r in rows:
        p = load_project(r["id"])
        est = price_project(p)
        plans = [s for s in p["sheets"] if s.get("kind") == "floor_plan"]
        sc = schedule(p.get("rate_schedule_id"))
        out.append(dict(id=p["id"], name=p["name"], status=p.get("status"), updated=p["updated"], total=est["total"],
                        created=p.get("created"), details=p.get("details") or dict(client=p.get("client", "")),
                        location=rates.loc_text(p.get("loc")) if p.get("loc") else (p.get("location") or "Not specified"),
                        rate_schedule=rates.label(sc) if sc else None,
                        files=len({s["file"] for s in p["sheets"]}), sheets=len(p["sheets"]),
                        plinth=round(plans[0]["metrics"]["plinth_area"], 1) if plans else 0))
    return out


@app.delete("/api/projects/{pid}/files/{fname}")
def remove_file(pid: str, fname: str):
    """Remove an uploaded drawing and every sheet read from it. Quantities are re-derived from the remaining
    drawings; the user's item edits (overrides) are kept — an edit on a line that no longer exists is simply unused."""
    p = load_project(pid)
    editable(p)
    fname = os.path.basename(fname)
    before = len(p["sheets"])
    p["sheets"] = [s for s in p["sheets"] if s.get("file") != fname]
    fp = os.path.join(UPLOADS, pid, fname)
    if before == len(p["sheets"]) and not os.path.exists(fp):
        raise HTTPException(404, "drawing not found in this project")
    if os.path.exists(fp):
        os.remove(fp)
    log(p, f"Removed drawing {fname} ({before - len(p['sheets'])} sheet(s))")
    refresh_derived(p)
    rebuild_items(p)
    save_project(p)
    return full_project(p)


@app.get("/api/scope-options")
def scope_options():
    return dict(types=scope_mod.TYPES, presets=scope_mod.PRESETS, categories=CATEGORY_ORDER, labels=scope_mod.LABELS,
                config_fields=scope_mod.CONFIG_FIELDS)


@app.get("/api/projects/{pid}/uploads")
def project_uploads(pid: str):
    """Files stored for this project (uploads/<pid>/), with size and the sheets read from each."""
    p = load_project(pid)
    folder = os.path.join(UPLOADS, pid)
    out = []
    for fn in sorted(os.listdir(folder)) if os.path.isdir(folder) else []:
        fp = os.path.join(folder, fn)
        if os.path.isfile(fp):
            out.append(dict(file=fn, bytes=os.path.getsize(fp), uploaded=datetime.fromtimestamp(os.path.getmtime(fp)).isoformat(timespec="seconds"),
                            sheets=sum(1 for sh in p["sheets"] if sh.get("file") == fn)))
    return out


MAX_UPLOAD = 200 * 1024 * 1024
MAGIC = {b"%PDF-": ".pdf", b"\x89PNG": ".png", b"\xff\xd8\xff": ".jpg", b"AC10": ".dwg", b"II*\x00": ".tif", b"MM\x00*": ".tif"}


def _validate_upload(name: str, data: bytes):
    """Size, magic bytes, PDF / image validity and image dimensions — nothing uploaded is ever executed."""
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, f"{name}: larger than 200 MB")
    ext = os.path.splitext(name.lower())[1]
    kind = next((v for k, v in MAGIC.items() if data.startswith(k)), None)
    if ext == ".dxf":
        if b"SECTION" not in data[:4096]:
            raise HTTPException(400, f"{name}: not a DXF file")
        return
    if kind is None or (kind != ext and not (kind == ".jpg" and ext == ".jpeg") and not (kind == ".tif" and ext == ".tiff")):
        raise HTTPException(400, f"{name}: file content does not match a supported drawing type (PDF, DXF, DWG, PNG, JPG, TIF)")
    if kind == ".pdf":
        import pymupdf
        try:
            pymupdf.open(stream=data, filetype="pdf").page_count
        except Exception:
            raise HTTPException(400, f"{name}: damaged / unreadable PDF")
    if kind in (".png", ".jpg", ".tif"):
        from PIL import Image
        try:
            w, h = Image.open(io.BytesIO(data)).size
        except Exception:
            raise HTTPException(400, f"{name}: unreadable image")
        if w * h > 400_000_000:
            raise HTTPException(413, f"{name}: image too large ({w}×{h})")


def _save_uploads(folder, files, skipped=None):
    """Save uploads into the project folder. A file identical (SHA-256) to one already in the project is NOT added
    again (it would double-count quantities) — it is reported in `skipped`. A different file with an existing name
    is stored under a new name instead of silently overwriting the earlier drawing."""
    paths = []
    os.makedirs(folder, exist_ok=True)
    have = {}
    for fn in os.listdir(folder):
        fp = os.path.join(folder, fn)
        if os.path.isfile(fp):
            have[hashlib.sha256(open(fp, "rb").read()).hexdigest()] = fn
    for f in files or []:
        if not f.filename:
            continue
        safe = re.sub(r"[^A-Za-z0-9._ -]", "_", os.path.basename(f.filename))
        data = f.file.read()
        _validate_upload(safe, data)
        h = hashlib.sha256(data).hexdigest()
        if h in have:
            if skipped is not None:
                skipped.append(dict(file=safe, reason=f"identical to '{have[h]}' already in this project — not added again"))
            continue
        stem, ext = os.path.splitext(safe)
        n, dst = 2, os.path.join(folder, safe)
        while os.path.exists(dst):
            safe = f"{stem} ({n}){ext}"
            dst = os.path.join(folder, safe)
            n += 1
        with open(dst, "wb") as fh:
            fh.write(data)
        have[h] = safe
        paths.append((dst, safe))
    return paths


@app.post("/api/projects")
def create_project(name: str = Form("Untitled estimate"), client: str = Form(""), location: str = Form(""),
                   params: str = Form("{}"), sample: str = Form(""), loc: str = Form("{}"), rate_schedule_id: str = Form(""),
                   building_type: str = Form(""), site_address: str = Form(""), description: str = Form(""),
                   scope: str = Form(""),
                   files: list[UploadFile] | None = File(None)):
    """Create a project. Drawings are optional here: a project can be created first and its drawings uploaded later
    with POST /api/projects/{pid}/files. Everything uploaded is stored under that project (DB record + uploads/<pid>/)."""
    try:
        loc_d = clean_loc(json.loads(loc or "{}"))
    except ValueError:
        raise HTTPException(400, "loc must be a JSON object")
    if rate_schedule_id:
        sc = schedule(rate_schedule_id)
        if not sc or sc["status"] != "active" or sc.get("region") != tn_scope.STATE:
            raise HTTPException(400, "Only the active Tamil Nadu PWD schedule can be used")
    pid = uuid.uuid4().hex[:10]
    folder = os.path.join(UPLOADS, pid)
    os.makedirs(folder, exist_ok=True)
    paths = []
    if sample:
        src = os.path.join(SAMPLES, os.path.basename(sample))
        if not os.path.exists(src):
            raise HTTPException(404, "sample not found")
        dst = os.path.join(folder, os.path.basename(sample))
        shutil.copy(src, dst)
        paths.append((dst, os.path.basename(sample)))
    paths += _save_uploads(folder, files)
    if not (name or "").strip():
        raise HTTPException(400, "Project name is required")
    sheets = []
    if paths:
        with LOCK:
            sheets = analyze_files(pid, paths)
    now = datetime.now().isoformat(timespec="seconds")
    p = dict(id=pid, name=name or "Untitled estimate", client=client, location=location or (rates.loc_text(loc_d) if loc_d else ""),
             created=now, updated=now, status="draft", params=dict(takeoff.DEFAULT_PARAMS, _src={}), sheets=sheets, items=[],
             overrides={}, history=[], versions=[], loc=loc_d,
             details=dict(client=client.strip(), building_type=building_type.strip(), site_address=site_address.strip(),
                          description=description.strip()))
    try:
        p["scope"] = scope_mod.normalize(json.loads(scope)) if scope else scope_mod.default()
    except (ValueError, TypeError):
        raise HTTPException(400, "scope must be a JSON object")
    with db() as con:
        res = tn_scope.resolve(con, loc_d)
    if rate_schedule_id:
        set_schedule(p, rate_schedule_id, "manual", "Chosen when the project was created", resolved=res)
    elif res["status"] == "selected":
        set_schedule(p, res["schedule"]["id"], "auto", res["reason"], resolved=res)
    else:
        set_schedule(p, None, "none", res["reason"], resolved=res)
    apply_scale_guard(p)
    refresh_derived(p)
    try:
        user = {k: v for k, v in json.loads(params or "{}").items() if v not in (None, "")}
    except Exception:
        user = {}
    for k, v in user.items():
        p["params"][k] = v
        p["params"]["_src"][k] = "user"
    log(p, f"Analysed {len(paths)} file(s) → {len(sheets)} sheet(s)" if paths else "Project created — awaiting drawings")
    rebuild_items(p)
    save_project(p)
    return full_project(p)


@app.post("/api/projects/{pid}/files")
def add_files(pid: str, files: list[UploadFile] = File(...)):
    p = load_project(pid)
    editable(p)
    skipped = []
    paths = _save_uploads(os.path.join(UPLOADS, pid), files, skipped)
    if paths:
        with LOCK:
            p["sheets"] += analyze_files(pid, paths, start=len(p["sheets"]))
        log(p, f"Added {len(paths)} drawing file(s)")
        apply_scale_guard(p)
    for k in skipped:
        log(p, f"Upload skipped: {k['file']} — {k['reason']}")
    refresh_derived(p)
    rebuild_items(p)  # item edits (overrides) are keyed by item and survive new drawings
    save_project(p)
    out = full_project(p)
    out["upload_skipped"] = skipped
    return out


@app.get("/api/projects/{pid}")
def get_project(pid: str):
    return full_project(load_project(pid))


@app.delete("/api/projects/{pid}")
def delete_project(pid: str):
    with db() as con:
        con.execute("DELETE FROM projects WHERE id=?", (pid,))
    shutil.rmtree(os.path.join(UPLOADS, os.path.basename(pid)), ignore_errors=True)
    shutil.rmtree(os.path.join(PREVIEWS, os.path.basename(pid)), ignore_errors=True)
    return dict(ok=True)


@app.patch("/api/projects/{pid}")
def update_project(pid: str, body: dict):
    p = load_project(pid)
    if "name" in body:
        if not str(body["name"]).strip():
            raise HTTPException(400, "Project name is required")
        p["name"] = str(body["name"]).strip()
    if "details" in body:  # client, building type, site address, description
        det = p.setdefault("details", {})
        for k in ("client", "building_type", "site_address", "description"):
            if k in body["details"]:
                det[k] = str(body["details"][k] or "").strip()
        p["client"] = det.get("client", p.get("client", ""))
        log(p, "Project details updated")
    if "scope" in body:  # changes what the estimate covers; quantities stay, totals are recomputed
        editable(p)
        old_area = ((p.get("scope") or {}).get("config") or {}).get("built_up_area_m2")
        p["scope"] = scope_mod.normalize(body["scope"])
        log(p, "Scope: " + scope_mod.describe(p["scope"]))
        if (p["scope"].get("config") or {}).get("built_up_area_m2") != old_area:  # new user area: re-check guessed scales
            for s_ in p["sheets"]:
                if (s_.get("scale_guard") or {}).get("applied") is False and not (s_.get("scale_guard") or {}).get("user_calibrated"):
                    s_.pop("scale_guard", None)
            apply_scale_guard(p)
        rebuild_items(p)
    if "params" in body:
        editable(p)
        srcs = p["params"].setdefault("_src", {})
        changed = []
        for k, v in body["params"].items():
            if k == "_src":
                continue
            if v is not None and k in PARAM_BOUNDS:
                lo, hi, cast = PARAM_BOUNDS[k]
                v = checked_number(k, v, lo, hi, cast)
            elif v is not None and k in CHOICES and v not in CHOICES[k]:
                raise HTTPException(400, f"{k} must be one of {', '.join(CHOICES[k])}")
            if v is None and k in DERIVABLE:  # back to the value read from the drawings
                srcs.pop(k, None)
                changed.append(f"{k}=from drawing")
                continue
            if p["params"].get(k) != v:
                changed.append(f"{k}={v}")
            p["params"][k] = v
            if k in DERIVABLE:
                srcs[k] = "user"
        refresh_derived(p)
        if changed:
            log(p, "Parameters: " + ", ".join(changed))
        rebuild_items(p)
    save_project(p)
    return full_project(p)


@app.post("/api/projects/{pid}/confirm")
def confirm(pid: str, body: dict):
    p = load_project(pid)
    editable(p)
    who = (body.get("by") or "").strip() or "Estimator"
    est = price_project(p)
    with db() as con:
        learning.record_confirmation(con, p, est)  # the estimator learns from every confirmed estimate
    sc = schedule(p.get("rate_schedule_id"))
    no = len(p.get("versions", [])) + 1
    p.setdefault("versions", []).append(dict(
        no=no, t=datetime.now().isoformat(timespec="seconds"), by=who, note=body.get("note", ""), total=est["total"], subtotal=est["subtotal"],
        location=rates.loc_text(p.get("loc")), rate_schedule_id=p.get("rate_schedule_id"), rate_schedule=rates.public(sc),
        items=[dict(sno=i["sno"], description=i["description"], qty=i["qty"], unit=i["unit"], rate=i["rate"], amount=i["amount"], sor_code=i.get("sor_code")) for i in est["items"]]))
    p["status"] = "confirmed"
    log(p, f"Revision R{no} confirmed and locked by {who} — total ₹{est['total']:,.0f}")
    save_project(p)
    return full_project(p)


@app.post("/api/projects/{pid}/reopen")
def reopen(pid: str, body: dict | None = None):
    p = load_project(pid)
    p["status"] = "in_review"
    log(p, f"Reopened for revision by {(body or {}).get('by') or 'Estimator'}")
    save_project(p)
    return full_project(p)


@app.get("/api/projects/{pid}/versions/{no}")
def get_version(pid: str, no: int):
    p = load_project(pid)
    v = next((x for x in p.get("versions", []) if x["no"] == no), None)
    if not v:
        raise HTTPException(404)
    return v


@app.patch("/api/projects/{pid}/items/{key}")
def update_item(pid: str, key: str, body: dict):
    p = load_project(pid)
    editable(p)
    if body.get("reset"):
        p["overrides"].pop(key, None)
        log(p, f"Reset item {key} to the drawing take-off")
    else:
        before = next((i for i in price_project(p)["items"] if i["key"] == key), None)
        o = p["overrides"].setdefault(key, {})
        if o.get("schedule_id") and o["schedule_id"] != p.get("rate_schedule_id"):
            for f in ("sor_code", "rate", "status"):  # choices made under another schedule no longer apply
                o.pop(f, None)
        for k, hi in (("qty", MAX_QTY), ("rate", MAX_RATE)):
            if body.get(k) is not None:
                body[k] = checked_number(k, body[k], 0.0, hi)  # no negative quantities / rates, no text
        for k in ("qty", "rate", "sor_code", "status", "note"):
            if k in body:
                if body[k] is None:
                    o.pop(k, None)
                else:
                    o[k] = body[k]
        if "sor_code" in body or "rate" in body or "status" in body:
            o["schedule_id"] = p.get("rate_schedule_id")
        if "sor_code" in body or "rate" in body:
            o.setdefault("changes", []).append(dict(
                at=datetime.now().isoformat(timespec="seconds"), schedule_id=p.get("rate_schedule_id"),
                from_code=before.get("sor_code") if before else None, to_code=body.get("sor_code", before.get("sor_code") if before else None),
                from_rate=before.get("base_rate") if before else None, to_rate=body.get("rate"), reason=body.get("reason", "")))
        if not {k for k in o if k not in ("changes", "schedule_id")}:
            p["overrides"].pop(key, None)
        log(p, f"Edited {key}: " + ", ".join(f"{k}={body[k]}" for k in body))
    save_project(p)
    return full_project(p)


@app.post("/api/projects/{pid}/items")
def add_item(pid: str, body: dict):
    p = load_project(pid)
    editable(p)
    key = "custom_" + uuid.uuid4().hex[:6]
    it = dict(key=key, description=body.get("description", "Custom item"), category=body.get("category", "Miscellaneous"),
              unit=body.get("unit", "sqm"), qty=checked_number("qty", body.get("qty", 1), 0.0, MAX_QTY), formula="Added manually by the estimator",
              sources=[], hint=None, extract_conf=1.0, drawing_refs=["manual entry"], custom=True, basis="manual")
    map_items(p, [it])
    p["items"].append(it)
    log(p, f"Added item: {it['description'][:60]}")
    save_project(p)
    return full_project(p)


@app.delete("/api/projects/{pid}/items/{key}")
def delete_item(pid: str, key: str):
    p = load_project(pid)
    editable(p)
    p["items"] = [i for i in p["items"] if not (i["key"] == key and i.get("custom"))]
    p["overrides"].pop(key, None)
    save_project(p)
    return full_project(p)


@app.post("/api/projects/{pid}/validate-all")
def validate_all(pid: str):
    p = load_project(pid)
    editable(p)
    est = price_project(p)
    for i in est["items"]:
        if i["status"] == "auto":
            o = p["overrides"].setdefault(i["key"], {})
            o["status"] = "confirmed"
            o["schedule_id"] = p.get("rate_schedule_id")
    log(p, "Validated all automatically mapped items")
    save_project(p)
    return full_project(p)


@app.patch("/api/projects/{pid}/sheets/{idx}/rooms/{rid}")
def update_room(pid: str, idx: int, rid: str, body: dict):
    p = load_project(pid)
    editable(p)
    for r in p["sheets"][idx]["rooms"]:
        if r["id"] == rid:
            if body.get("name"):
                r["name"] = body["name"]
            if body.get("type"):
                r["type"], r["type_conf"], r["type_source"] = body["type"], 1.0, "user"
            log(p, f"Room {rid} → {r['name']} ({r['type']})")
    rebuild_items(p)
    save_project(p)
    return full_project(p)


@app.patch("/api/projects/{pid}/sheets/{idx}")
def update_sheet(pid: str, idx: int, body: dict):
    """Use/skip a sheet, set its type or floor, or re-run analysis with a calibrated scale."""
    p = load_project(pid)
    editable(p)
    if not 0 <= idx < len(p["sheets"]):
        raise HTTPException(404, "sheet not found in this project")
    s = p["sheets"][idx]
    if "enabled" in body:
        s["enabled"] = bool(body["enabled"])
    if body.get("kind"):
        s["kind"] = body["kind"]
    if "scale_confirmed" in body:  # user accepts the automatic scale as a stated assumption (still flagged as not measured)
        s["scale_confirmed"] = bool(body["scale_confirmed"])
        log(p, f"{s['file']} {s.get('sheet', '')}: automatic scale {'accepted by user as an assumption' if s['scale_confirmed'] else 'acceptance withdrawn'}")
    if "opening_confirmed" in body:  # user-confirmed door / window counts — replace the automatic count everywhere
        oc = body["opening_confirmed"] or {}
        conf = {k: int(checked_number(f"{k} count", oc[k], 0, 5000)) for k in ("door", "window") if oc.get(k) not in (None, "")}
        if conf:
            s["opening_confirmed"] = conf
        else:
            s.pop("opening_confirmed", None)
        log(p, f"{s['file']} {s.get('sheet', '')}: opening count confirmed by user {conf or '(cleared)'}")
    if "label" in body:  # user label (Bedroom, Ground floor, Electrical layout, custom…) — organisation only
        s["user_label"] = str(body["label"] or "").strip()[:80]
        log(p, f"{s['file']} {s.get('sheet', '')} labelled '{s['user_label'] or '—'}'")
    if "floor" in body:
        s["floor"] = body["floor"]
        s["floor_confirmed"] = body["floor"] is not None  # floors set by the user (vs read from drawing text)
        log(p, f"{s['file']} {s.get('sheet', '')} set as {body['floor'].get('label')} (×{body['floor'].get('repeat', 1)})")
    if body.get("calib"):
        c = body["calib"]
        real = checked_number("real_m", c.get("real_m"), 0.01, 1000.0)
        meas = checked_number("measured_m", c.get("measured_m"), 0.01, 100000.0)
        k = real / meas
        opts = {}
        if s.get("source") == "pdf":
            opts["scale_den"] = int(round(s["scale"].get("scale_den", 100) * k))
        elif s.get("source") == "image":
            opts["px_per_m"] = s["scale"]["px_per_m"] / k
        page0 = _rescale_sheet(p, idx, k, "calibrated by user") if opts else None
        if page0 is not None:
            page0["scale"]["confidence"] = 0.98
            page0["scale_guard"] = dict(applied=False, user_calibrated=True)
        log(p, f"Recalibrated {s['file']} {s.get('sheet', '')}")
    refresh_derived(p)
    rebuild_items(p)
    save_project(p)
    return full_project(p)


# ============================================================================ jurisdiction + rate schedules
# regions offered even before their official schedule is loaded (the resolver then says so explicitly)
ALWAYS_LIST = {"Goa"}
LOC_KEYS = ("country", "region", "district", "city", "municipality", "authority_level")


def clean_loc(d):
    if not isinstance(d, dict):
        raise ValueError
    out = {k: str(d[k]).strip() for k in LOC_KEYS + ("executing_authority",) if d.get(k) not in (None, "") and str(d[k]).strip()}
    try:
        return tn_scope.tn_loc(out)
    except tn_scope.OutOfScope as e:
        raise HTTPException(400, str(e))


@app.get("/api/locations")
def locations():
    return tn_scope.locations()


def _legacy_locations():
    ref = rates.load_regions(JURIS)
    with db() as con:
        sch = rates.list_schedules(con)
    # countries / regions / districts / cities used by configured schedules are offered too
    names = {c["name"].lower() for c in ref["countries"]}
    for s in sch:
        if s.get("country") and s["country"].lower() not in names:
            ref["countries"].append(dict(name=s["country"], region_label="Region", district_label="District", city_label="City", regions=[]))
            names.add(s["country"].lower())
    for c in ref["countries"]:
        have = {r["name"].lower() for r in c["regions"]}
        mine = [s for s in sch if (s.get("country") or "").lower() == c["name"].lower()]
        for s in mine:
            if s.get("region") and s["region"].lower() not in have:
                c["regions"].append(dict(name=s["region"], type="Region"))
                have.add(s["region"].lower())
        c["districts"] = sorted({s["district"] for s in mine if s.get("district")})
        c["cities"] = sorted({s["city"] for s in mine if s.get("city")})
    # offer only the states / regions for which an official schedule is actually loaded and active
    active = {(s.get("country") or "").lower(): set() for s in sch}
    for s in sch:
        if s.get("status") == "active" and s.get("region"):
            active.setdefault((s.get("country") or "").lower(), set()).add(s["region"].lower())
    for c in ref["countries"]:
        c["regions"] = [r for r in c["regions"] if r["name"].lower() in active.get(c["name"].lower(), set())
                        or r["name"] in ALWAYS_LIST]
        c["districts"] = sorted({s["district"] for s in sch if s.get("status") == "active" and s.get("district")
                                 and (s.get("country") or "").lower() == c["name"].lower()})
    ref["authority_levels"] = [dict(value=v, label=l) for v, l in rates.AUTHORITY_LEVELS]
    ref.pop("_note", None)
    return ref


@app.get("/api/rate-schedules")
def get_schedules(country: str = "", region: str = "", authority: str = "", authority_level: str = "", year: str = "", status: str = ""):
    with db() as con:
        return [rates.public(s) for s in rates.list_schedules(con, country=country, region=region, authority=authority,
                                                              authority_level=authority_level, year=year, status=status)]


@app.get("/api/rate-schedules/applicable")
def applicable_schedule(country: str = "", region: str = "", district: str = "", city: str = "", authority_level: str = "", on: str = ""):
    loc = clean_loc(dict(country=country, region=region, district=district, city=city, authority_level=authority_level))
    with db() as con:
        r = tn_scope.resolve(con, loc, on or None)
    r["location_text"] = rates.loc_text(loc)
    return r


@app.get("/api/rate-schedules/{sid}")
def get_schedule_ep(sid: str):
    with db() as con:
        s = rates.get_schedule(con, sid)
    if not s:
        raise HTTPException(404, "Rate schedule not found")
    return rates.public(s)


@app.post("/api/rate-schedules")
def create_schedule_ep(body: dict):
    with db() as con:
        sid, errs = rates.create_schedule(con, body, by=body.get("by") or "administrator")
    if errs:
        raise HTTPException(422, dict(message="Schedule not created", errors=errs))
    return get_schedule_ep(sid)


IDENTITY = ("country", "region", "district", "city", "authority", "authority_level", "schedule_type", "code_system", "year",
            "financial_year", "revision")


@app.patch("/api/rate-schedules/{sid}")
def update_schedule_ep(sid: str, body: dict):
    with db() as con:
        s = rates.get_schedule(con, sid)
        if not s:
            raise HTTPException(404, "Rate schedule not found")
        meta = {k: body[k] for k in rates.META_FIELDS if k in body}
        if meta:
            if s["status"] == "active" and any(k in meta for k in IDENTITY):
                raise HTTPException(409, "Deactivate the schedule before changing its identity or jurisdiction")
            errs = rates.validate_meta(con, {**s, **meta}, sid)
            if errs:
                raise HTTPException(422, dict(message="Invalid schedule metadata", errors=errs))
            con.execute(f"UPDATE rate_schedules SET {', '.join(k + '=?' for k in meta)} WHERE id=?",
                        (*[rates._n(v) for v in meta.values()], sid))
        st = body.get("status")
        if st:
            if st not in rates.STATUSES:
                raise HTTPException(422, "status must be draft, active or inactive")
            if st == "active" and s["status"] != "active":
                errs = rates.activation_errors(con, sid)
                if errs:
                    raise HTTPException(409, dict(message="Schedule not activated", errors=errs))
            con.execute("UPDATE rate_schedules SET status=? WHERE id=?", (st, sid))
    LIB.invalidate(sid)
    return get_schedule_ep(sid)


def _read_table(file: UploadFile):
    data = file.file.read()
    if file.filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            header = None
            for vals in ws.iter_rows(values_only=True):
                cells = [str(v).strip().lower() if v is not None else "" for v in vals]
                if header is None:
                    if any(c in ("description", "desc", "item description") for c in cells) and any("rate" in c for c in cells):
                        header = cells
                    continue
                rows.append(dict(zip(header, vals)))
        return rows
    return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig", errors="ignore"))))


def _import(sid, file, by, replace):
    rows = _read_table(file)
    with db() as con:
        r = rates.import_items(con, sid, rows, by=by or "administrator", replace=replace)
    if not r["ok"]:
        raise HTTPException(422, dict(message="Import rejected — nothing was changed", **{k: v for k, v in r.items() if k != "ok"}))
    LIB.invalidate(sid)
    return r


@app.post("/api/rate-schedules/{sid}/import")
def import_schedule_items(sid: str, file: UploadFile = File(...), by: str = Form(""), replace: bool = Form(False)):
    return _import(sid, file, by, replace)


@app.patch("/api/projects/{pid}/location")
def set_location(pid: str, body: dict):
    p = load_project(pid)
    editable(p)  # a confirmed estimate must be reopened (new revision) first
    try:
        loc = clean_loc(body.get("loc") or {})
    except ValueError:
        raise HTTPException(400, "loc must be an object")
    old_loc = rates.loc_text(p.get("loc"))
    p["loc"] = loc
    p["location"] = rates.loc_text(loc) if loc else ""
    log(p, f"Location: {old_loc} → {rates.loc_text(loc)}")
    with db() as con:
        res = tn_scope.resolve(con, loc)
    cur = p.get("rate_schedule_id")
    if res["status"] == "selected":
        if res["schedule"]["id"] != cur:
            set_schedule(p, res["schedule"]["id"], "auto", res["reason"], by=body.get("by"), resolved=res)
            notice = "The applicable rate schedule changed, so PWD/SOR rates and the estimated cost were recalculated. Quantities are unchanged."
        else:
            notice = "The rate schedule is unchanged: " + res["reason"]
    elif cur and any(c["id"] == cur for c in res["candidates"]):
        notice = "The current rate schedule still applies to the new location. " + res["reason"]
    elif cur:
        set_schedule(p, None, "none", res["reason"], by=body.get("by"), resolved=res)
        notice = ("The previous rate schedule does not apply to the new location and was removed. " + res["reason"]
                  + " Quantities are unchanged; items have no rate until a schedule is selected.")
    else:
        p["rate_selection"] = dict(p.get("rate_selection") or {}, status=res["status"], reason=res["reason"])
        notice = res["reason"]
    rebuild_items(p)
    save_project(p)
    out = full_project(p)
    out["resolution"] = res
    out["notice"] = notice
    return out


@app.patch("/api/projects/{pid}/rate-schedule")
def set_rate_schedule(pid: str, body: dict):
    p = load_project(pid)
    editable(p)
    with db() as con:
        res = tn_scope.resolve(con, p.get("loc"))
    if body.get("auto"):
        if res["status"] != "selected":
            raise HTTPException(409, dict(message=res["reason"], candidates=res["candidates"]))
        set_schedule(p, res["schedule"]["id"], "auto", res["reason"], by=body.get("by"), resolved=res)
    else:
        sid = body.get("schedule_id")
        sc = schedule(sid) if sid else None
        if not sc:
            raise HTTPException(404, "Rate schedule not found")
        if sc["status"] != "active" or sc.get("region") != tn_scope.STATE:
            raise HTTPException(409, "Only the active Tamil Nadu PWD schedule can be applied to a project")
        applicable = any(c["id"] == sid for c in res["candidates"])
        reason = (body.get("reason") or "").strip()
        if not applicable and not reason:
            raise HTTPException(409, dict(message="This schedule is not configured for the project's jurisdiction. "
                                                  "Give a reason to apply it anyway (recorded in the audit trail).",
                                          applicable=[c["id"] for c in res["candidates"]]))
        set_schedule(p, sid, "manual", reason or "Selected by the user from the applicable schedules", by=body.get("by"), resolved=res)
        p["rate_selection"]["outside_jurisdiction"] = not applicable
    rebuild_items(p)
    save_project(p)
    return full_project(p)


# ============================================================================ rate library (items of one schedule) + rate analysis
def _sid(schedule_id: str = ""):
    sid = schedule_id or rates.CPWD_ID
    if not schedule(sid):
        raise HTTPException(404, "Rate schedule not found")
    return sid


@app.get("/api/sor")
def list_sor(q: str = "", category: str = "", schedule_id: str = ""):
    sid = _sid(schedule_id)
    with db() as con:
        rows = rates.items_of(con, sid)
    if category:
        rows = [r for r in rows if r["category"] == category]
    if q:
        rows = [r for r in rows if q.lower() in (r["description"] + r["code"]).lower()]
    rows.sort(key=lambda r: (CATEGORY_ORDER.index(r["category"]) if r["category"] in CATEGORY_ORDER else 99, r["code"]))
    return rows


@app.get("/api/sor/search")
def search_sor(q: str, unit: str = "", k: int = 8, schedule_id: str = ""):
    _, matcher = library(_sid(schedule_id))
    return clean(matcher.search(q, unit=unit or None, k=k)) if matcher else []


@app.put("/api/sor/{code}")
def upsert_sor(code: str, body: dict, schedule_id: str = ""):
    sid = _sid(schedule_id or body.get("schedule_id", ""))
    with db() as con:
        cur = con.execute("SELECT * FROM rate_items WHERE schedule_id=? AND code=?", (sid, code)).fetchone()
        d = dict(cur) if cur else dict(code=code, description="", unit="sqm", rate=0, category="Miscellaneous", source="user", analysis=None)
        d.update({k: body[k] for k in ("description", "unit", "rate", "category") if k in body})
        src = "user" if cur is None else ("edited" if d.get("source") in ("seed", "DSR 2023", "import") else d.get("source"))
        if "analysis" in body:
            a = body["analysis"]
            d["analysis"] = json.dumps(a) if a else None
            if a:
                d["rate"] = analysis_rate(a)
                src = "analysis"
        try:
            rate = float(d["rate"])
        except (TypeError, ValueError):
            raise HTTPException(422, "rate must be a number")
        con.execute("INSERT OR REPLACE INTO rate_items VALUES(?,?,?,?,?,?,?,?,?)",
                    (sid, code, d["description"], d["unit"], rate, d["category"], src, datetime.now().isoformat(timespec="seconds"),
                     d.get("analysis")))
    LIB.invalidate(sid)
    return dict(ok=True, rate=rate, schedule_id=sid)


@app.delete("/api/sor/{code}")
def delete_sor(code: str, schedule_id: str = ""):
    sid = _sid(schedule_id)
    with db() as con:
        con.execute("DELETE FROM rate_items WHERE schedule_id=? AND code=?", (sid, code))
    LIB.invalidate(sid)
    return dict(ok=True)


@app.post("/api/sor/reset")
def reset_sor():
    """Restore CPWD DSR 2023 from its source file (other schedules are not touched)."""
    with db() as con:
        rates.seed_cpwd(con, DSR_CSV, _seed_fallback)
    LIB.invalidate(rates.CPWD_ID)
    return dict(ok=True)


@app.post("/api/sor/import")
def import_sor(file: UploadFile = File(...), schedule_id: str = Form(""), by: str = Form("")):
    """Kept for compatibility: same validation as /api/rate-schedules/{id}/import."""
    return _import(_sid(schedule_id), file, by, False)


@app.get("/api/sor/export.csv")
def export_sor(schedule_id: str = ""):
    sid = _sid(schedule_id)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["code", "description", "unit", "rate", "category"])
    for r in list_sor(schedule_id=sid):
        w.writerow([r["code"], r["description"], r["unit"], r["rate"], r["category"]])
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv", headers={"Content-Disposition": f"attachment; filename={sid}.csv"})


# ============================================================================ exports
def _fname(full, ext):
    return re.sub(r"[^A-Za-z0-9]+", "_", full["name"]).strip("_") + f"_estimate.{ext}"


@app.get("/api/projects/{pid}/export.xlsx")
def export_xlsx(pid: str):
    from reports import build_xlsx
    full = full_project(load_project(pid))
    return StreamingResponse(io.BytesIO(build_xlsx(full)), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={_fname(full, 'xlsx')}"})


@app.get("/api/projects/{pid}/report.pdf")
def report_pdf(pid: str):
    from reports import build_pdf
    full = full_project(load_project(pid))
    return Response(build_pdf(full, PREVIEWS), media_type="application/pdf",
                    headers={"Content-Disposition": f"inline; filename={_fname(full, 'pdf')}"})


@app.get("/api/projects/{pid}/report", response_class=HTMLResponse)
def report(pid: str):
    from reports import build_html
    return build_html(full_project(load_project(pid)))


# ============================================================================ frontend
app.mount("/static", StaticFiles(directory=FRONT), name="static")


@app.get("/", response_class=HTMLResponse)
def index():
    with open(os.path.join(FRONT, "index.html"), encoding="utf-8") as f:
        return f.read()


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=int(os.environ.get("PORT", 8000)))
